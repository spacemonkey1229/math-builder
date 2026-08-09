#!/usr/bin/env python3
"""
submission-score — MACHINE FIRST PASS against RUBRIC.md.

    python3 skills/submission-score/scripts/score_submission.py \
        --learner avery-n --case cold-chain \
        --file /path/to/pitch.docx [--file /path/to/workbook.xlsx] \
        [--project .] [--write]

Produces a COMPLETE first-pass verdict on all seven criteria — that is the
scalability requirement: a guide reviewing twenty submissions starts from a
filled-in form, not a blank one.

THREE HARD LIMITS (RUBRIC.md, binding — enforced in code, not just documented):
  1. A machine BLOCKER on C3 is ADVISORY, not final.  Recomputation disagreements
     are emitted as "check this", never as "wrong".  See _c3().
  2. The machine may NEVER mark C7 Exemplary.  Enforced by a hard clamp in _c7()
     and re-asserted before write.
  3. Every verdict is `machine-first-pass` until a guide confirms.  This script
     cannot write status "confirmed"; only guide_confirm.py can.

Two outputs, never averaged:
  * mastery gate  badge_eligible = all(criterion >= Proficient)
  * growth score  0-100 weighted; Not Yet 0 / Proficient 70 / Exemplary 100
"""
import argparse, datetime, json, os, re, shutil, sys


def snapshot(project, filenames):
    """MCS-22 -- nothing is superseded silently. Same pattern as case-reconcile's
    snapshot(): copy the frozen file to readmodel/_snapshots/<file>.<timestamp>
    before this script overwrites it."""
    ts = datetime.datetime.now().strftime('%Y-%m-%d_%H%M%S')
    snap_dir = os.path.join(project, 'readmodel', '_snapshots')
    os.makedirs(snap_dir, exist_ok=True)
    for fn in filenames:
        src = os.path.join(project, 'readmodel', fn)
        if os.path.exists(src):
            shutil.copy2(src, os.path.join(snap_dir, f"{fn}.{ts}"))

LEVELS = {"Not Yet": 0, "Proficient": 70, "Exemplary": 100}
ORDER = ["Not Yet", "Proficient", "Exemplary"]
WEIGHTS = {"C1": 5, "C2": 20, "C3": 20, "C4": 15, "C5": 15, "C6": 15, "C7": 10}
NAMES = {"C1": "Research — the 'how to' resource", "C2": "The pitch — a firm stance, defended",
         "C3": "Mathematical correctness", "C4": "Required deliverables",
         "C5": "Visual display and appendix of work", "C6": "Socratic participation",
         "C7": "Framing the question"}
# RUBRIC.md tier adjustment table -> the criterion each requirement gates.
TIER_GATE = {
    "T1": None,
    "T2": ("C2", "the trade-off between the two competing variables is named explicitly",
           r"(?i)trade[\s-]?off|at the cost of|in exchange for|you give up|the price of (?:this|that)"),
    "T3": ("C3", "some case data is correctly identified as irrelevant and set aside",
           r"(?i)\b(ignor\w+|irrelevant|set (?:it |them )?aside|did not use|didn'?t use|"
           r"no bearing|not needed|red herring|left out)\b"),
    "T4": ("C3", "at least one assumption stated, justified, and its sensitivity tested",
           r"(?i)\b(if .{0,40}(?:instead|were|changed)|sensitiv\w+|i assume\w*|assuming\b)"),
}
NUM_RE = re.compile(r'(?<![\w.])(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\d+)(?![\w])')


# ------------------------------------------------------------------ intake --
def case_glossary_urls(cdir):
    """MCS-39 support for C1. Every case's published interactive HTML ships a
    hover+hyperlink glossary that already answers the exact "how to..." questions
    Challenge 1 asks the Hero to research independently. Extract every URL that
    glossary embeds so _c1() can tell "the Hero found this" from "the case handed
    this to them" -- same two authoring styles case-readiness's T1-19 recognizes
    (static <a class="gterm" href="..."> anchors, or a JS {label: {def, url}}
    data object rendered at load time). Returns a set of URLs; empty if the case
    ships no interactive HTML yet.
    """
    pub = os.path.join(cdir, 'published')
    urls = set()
    if not os.path.isdir(pub):
        return urls
    for fn in os.listdir(pub):
        if not fn.lower().endswith('.html'):
            continue
        txt = open(os.path.join(pub, fn), encoding='utf-8', errors='ignore').read()
        data_entries = re.findall(
            r'"[A-Za-z][A-Za-z \-]*":\s*\{def:\s*"(?:[^"\\]|\\.)*"\s*,\s*url:\s*"(https://[^"]+)"',
            txt)
        static_terms = [u for u in re.findall(r'class="gterm"[^>]*href="([^"]+)"', txt)
                         if '${' not in u]
        urls.update(data_entries if data_entries else static_terms)
    return urls


def numbers_in(t):
    out = set()
    for m in NUM_RE.finditer(t or ''):
        try:
            out.add(float(m.group(1).replace(',', '')))
        except ValueError:
            pass
    return out


def close(a, b, rel=0.005):
    """0.5% RELATIVE tolerance — a Hero who rounds 8571.43 to 8571 has not made an
    error.  No absolute floor: a floor of 1.0 makes 0.1046 'agree' with 0.1188,
    which is how a scorer silently passes wrong work."""
    if a == b:
        return True
    scale = max(abs(a), abs(b))
    return scale > 0 and abs(a - b) <= rel * scale


def ingest(paths):
    """-> {text, tables[], images:int, charts:int, chart_numbers:set, sources:[...]}"""
    doc = {"text": "", "tables": [], "images": 0, "charts": 0,
           "chart_numbers": set(), "read": [], "unreadable": []}
    for p in paths:
        ext = os.path.splitext(p)[1].lower()
        try:
            if ext in ('.md', '.txt'):
                t = open(p, encoding='utf-8').read()
                doc["text"] += "\n" + t
                doc["images"] += len(re.findall(r'!\[[^\]]*\]\([^)]+\)', t))
                for blk in re.findall(r'(?ms)^\|.*?(?:\n\s*\n|\Z)', t):
                    doc["tables"].append(blk)
            elif ext == '.docx':
                import docx
                d = docx.Document(p)
                doc["text"] += "\n" + "\n".join(par.text for par in d.paragraphs)
                for tb in d.tables:
                    rows = [" | ".join(c.text for c in r.cells) for r in tb.rows]
                    doc["tables"].append("\n".join(rows))
                    doc["text"] += "\n" + "\n".join(rows)
                doc["images"] += sum(1 for r in d.part.rels.values() if "image" in r.reltype)
            elif ext == '.xlsx':
                import openpyxl
                wb = openpyxl.load_workbook(p, data_only=True)
                for ws in wb.worksheets:
                    doc["charts"] += len(getattr(ws, '_charts', []))
                    for ch in getattr(ws, '_charts', []):
                        for s in getattr(ch, 'series', []):
                            for ref in (getattr(getattr(s, 'val', None), 'numRef', None),
                                        getattr(getattr(s, 'cat', None), 'numRef', None)):
                                if ref is None or not ref.f:
                                    continue
                                doc["chart_numbers"] |= _range_values(wb, ref.f)
                    rows = []
                    for r in ws.iter_rows(values_only=True):
                        if any(v is not None for v in r):
                            rows.append(" | ".join('' if v is None else str(v) for v in r))
                    if rows:
                        doc["tables"].append("\n".join(rows))
                        doc["text"] += "\n" + "\n".join(rows)
            elif ext == '.pdf':
                import pdfplumber
                with pdfplumber.open(p) as pdf:
                    for pg in pdf.pages:
                        doc["text"] += "\n" + (pg.extract_text() or "")
                        doc["images"] += len(pg.images)
            else:
                doc["unreadable"].append(f"{os.path.basename(p)} (unsupported extension {ext})")
                continue
            doc["read"].append(os.path.basename(p))
        except Exception as e:                       # never guess at unreadable content
            doc["unreadable"].append(f"{os.path.basename(p)}: {e!r}")
    return doc


def _range_values(wb, f):
    m = re.match(r"^'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$", f)
    if not m:
        return set()
    ws = wb[m.group(1)] if m.group(1) in wb.sheetnames else None
    if ws is None:
        return set()
    out = set()
    for row in ws[f"{m.group(2)}{m.group(3)}:{m.group(4)}{m.group(5)}"]:
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                out.add(float(c.value))
    return out


# ------------------------------------------------- the case's value table ---
_UNIT_TOKENS = {'usd', 'pct', 'ft', 'cm3', 'cm2', 'sqft', 'h', 'hr', 'hrs', 'mb',
                'm', 'kg', 'l', 'lbs', 'doses', 'min', 'mph'}


def _derived_aliases(key):
    """UAT FINDING 2026-08-08 (see student_work/UAT_LOG.md): value_table() below
    originally only recognised Cold Chain's own derived.containers[] shape, so
    C3 recomputation silently found ZERO checkable quantities for every other
    case in the corpus (each case stores derived{} as flat, per-case-named
    key->number pairs, not containers[]) -- every non-Cold-Chain submission was
    landing on the 'no recomputable quantity was found' Not Yet branch
    regardless of correctness. This generates best-effort natural-language
    aliases for an arbitrary flat derived{} key so a Hero's own prose has a
    real chance of matching it: the full underscore-to-space label, the same
    label with a trailing unit-ish token stripped (a Hero writes 'the margin
    is 0.634', not 'the margin ft is 0.634'), and a short tail of the last two
    tokens, mirroring the alias pattern the containers[] branch already used
    (aliases = [n] + [toks[-1]]). This is a first-pass heuristic, not a
    guarantee -- watch how it performs against real submissions and refine it,
    the same 'first cohort' caveat RUBRIC.md applies to every other untested
    part of the machine pass."""
    toks = key.split('_')
    aliases = {" ".join(toks)}
    if len(toks) > 1 and toks[-1] in _UNIT_TOKENS:
        aliases.add(" ".join(toks[:-1]))
    if len(toks) > 2:
        aliases.add(" ".join(toks[-2:]))
    return [a for a in aliases if a]


def value_table(case_md, c):
    """Every quantity the machine can independently recompute or look up, with a
    label to find it by.  This is what 'recompute the Hero's numbers' means."""
    T = []
    der = c.get('derived') or {}
    for row in (der.get('containers') or []):
        n = row.get('name', '')
        toks = n.split()
        aliases = [n] + ([toks[-1]] if len(toks) > 1 else [])
        for field, unit in (('volume_cm3', 'cm3'), ('surface_area_cm2', 'cm2'),
                            ('sa_v', 'ratio'), ('cold_life_h', 'h')):
            if row.get(field) is not None:
                T.append({"label": n, "field": field, "value": float(row[field]),
                          "unit": unit, "src": "case.json.derived", "aliases": aliases})
    if not der.get('containers'):
        # Generic case (every case in the corpus except Cold Chain) -- derived{}
        # is a flat dict of case-specific numeric fields. See _derived_aliases().
        for k, val in der.items():
            if k.startswith('_') or isinstance(val, bool) or not isinstance(val, (int, float)):
                continue
            T.append({"label": k.replace('_', ' '), "field": k, "value": float(val),
                      "unit": "", "src": "case.json.derived", "aliases": _derived_aliases(k)})
    bal = c.get('balance') or c.get('balance_test') or {}
    ALIAS = {"A_in_spec_doses": ["in spec", "in-spec", "within spec"],
             "B_in_spec_doses": ["in spec", "in-spec", "within spec"],
             "margin_pct": ["margin", "wins by", "ahead by", "percentage point"]}
    for k in ('A_in_spec_doses', 'B_in_spec_doses', 'margin_pct'):
        if bal.get(k) is not None:
            T.append({"label": k.replace('_', ' '), "field": k, "value": float(bal[k]),
                      "unit": "", "src": "case.json.balance_test",
                      "aliases": ALIAS.get(k, [k.replace('_', ' ')])})
    return T


def recompute(doc_text, table):
    """For each known quantity, find the Hero's figure near its label and compare.
    Emits AGREES / CHECK / NOT ADDRESSED.  Never 'WRONG' — hard limit 1.

    Two deliberate scoping choices, both learned from false positives:
      * the search window is the LINE containing the label, not a character
        radius.  A radius spills across spreadsheet rows and "finds" the next
        container's cold life.  A paragraph and a table row are both one line.
      * the candidate reported as the Hero's figure is the one nearest IN THE
        TEXT, not the one nearest in value.  "a margin of 17.1%" should report
        17.1, not whichever number in the sentence happens to sit closest to
        the expected 11.0."""
    out = []
    low = doc_text.lower()
    # line index: offset -> (line_start, line_end)
    bounds = []
    pos = 0
    for line in doc_text.split('\n'):
        bounds.append((pos, pos + len(line)))
        pos += len(line) + 1

    def line_of(off):
        for a, b in bounds:
            if a <= off <= b:
                return a, b
        return max(0, off - 80), off + 80

    for e in table:
        cands = []                      # (text_distance, value)
        for alias in (e.get('aliases') or [e['label']]):
            for m in re.finditer(re.escape(alias.lower()), low):
                a, b = line_of(m.start())
                for nm in NUM_RE.finditer(doc_text[a:b]):
                    try:
                        v = float(nm.group(1).replace(',', ''))
                    except ValueError:
                        continue
                    cands.append((abs((a + nm.start()) - m.start()), v))
        if not cands:
            out.append({**e, "status": "NOT ADDRESSED"})
            continue
        hit = next((v for _, v in sorted(cands) if close(v, e['value'])), None)
        if hit is not None:
            out.append({**e, "status": "AGREES", "hero": hit})
            continue
        # Only a candidate of PLAUSIBLY THE SAME QUANTITY can be a disagreement.
        # "a margin of 17.1%" sits between "7000" and "17.1"; 7000 is not a rival
        # reading of an 11.0% margin, it is a different quantity in the same
        # sentence.  Take the nearest candidate WITHIN AN ORDER OF MAGNITUDE, and
        # if there is none, the Hero simply never stated this figure —
        # NOT ADDRESSED, not a defect.
        plaus = [v for _, v in sorted(cands)
                 if e['value'] != 0 and 0.2 <= abs(v / e['value']) <= 5.0]
        if plaus:
            out.append({**e, "status": "CHECK", "hero": plaus[0],
                        "note": f"machine has {e['value']}, the Hero's figure next to this label "
                                f"is {plaus[0]}. The Hero may be right — check before acting."})
        else:
            out.append({**e, "status": "NOT ADDRESSED"})
    return out


# ------------------------------------------------------------- criteria -----
def _lvl(x):
    return {"level": x[0], "confidence": x[1], "evidence": x[2]}


def _c1(doc, c, sub, glossary_urls=None):
    glossary_urls = glossary_urls or set()
    url = (sub or {}).get('challenge_1_resource_url') or ''
    urls = re.findall(r'https?://\S+', doc["text"]) + ([url] if url else [])
    # \S+ over-captures trailing punctuation a sentence would leave attached
    # ("...value, it explains" -> "...value,"); strip it before any comparison
    # or the MCS-39 glossary-repost check below silently never matches.
    urls = [u.rstrip('.,;:!?)]}\'"') for u in urls]
    if not urls:
        return _lvl(("Not Yet", "HIGH", ["no resource link found in the submission or the tracker"]))
    # MCS-39 — every case now ships a built-in glossary that already answers the
    # case's "how to..." questions. Reposting one of those URLs verbatim isn't
    # evidence of independent research; strip a trailing slash before comparing
    # since that's the only normalization worth doing without fetching anything.
    reposted = [u for u in urls if u.rstrip('/') in {g.rstrip('/') for g in glossary_urls}]
    labels = [d for d in re.findall(r'(?i)how to [a-z \-]+', doc["text"])]
    ev = [f"{len(urls)} link(s) found, first: {urls[0][:70]}"]
    if reposted:
        ev.append(f"MCS-39 — {len(reposted)} of the submitted link(s) are byte-identical to a "
                  f"URL already embedded in this case's own interactive-HTML glossary: "
                  f"{reposted[0][:70]}. This does not by itself demonstrate the Hero found and "
                  f"used the resource independently (RUBRIC.md C1 note) — a guide should look "
                  f"closer before confirming Proficient.")
    if labels:
        ev.append(f"names a 'how to…' skill in prose: {labels[:2]}")
        learned = re.search(r"(?i)i (?:got it wrong|had it wrong|thought .{0,40}then)|"
                            r"(?:got|had) it wrong (?:twice|once|at first)|before .{0,30}stuck|"
                            r"i did not know how|i didn'?t know how|my first (?:wrong )?attempt|"
                            r"false start|then i found", doc["text"])
        if learned and not reposted:
            ev.append(f"shows the skill being LEARNED, not just cited: {learned.group(0)[:60]!r}")
            return _lvl(("Exemplary", "LOW", ev + [
                "Exemplary requires the resource to address a skill the Hero genuinely did not "
                "have. Cue-phrase detection only — a guide must confirm it is real and not a "
                "rhetorical flourish.",
                "link liveness and topical match to Appendix G NOT verified — the machine cannot fetch"]))
        if learned and reposted:
            # Learned-language present alongside a reposted glossary URL is a genuine
            # ambiguous case, not a clean Exemplary or a clean Not Yet -- surface it,
            # don't guess past it.
            ev.append(f"ALSO shows learned-language cues: {learned.group(0)[:60]!r} — ambiguous "
                      f"given the reposted glossary link; guide call, not a machine one")
            return _lvl(("Proficient", "LOW", ev))
        return _lvl(("Proficient", "MEDIUM" if not reposted else "LOW", ev + [
            "no evidence of the skill being LEARNED (a false start, a correction) — not Exemplary",
            "link liveness and topical match to Appendix G NOT verified — the machine cannot fetch"]))
    return _lvl(("Proficient", "LOW", ev + [
        "no 'how to…' phrasing in the text, so the link's relevance to Appendix G is unverified",
        "Exemplary requires visible evidence of the skill being LEARNED (a false start, a "
        "correction) — a guide must read for that"]))


def _c2(doc, c, tier_note):
    text = doc["text"]
    # word count excludes the appendix of work per Challenge 4
    prose = text
    for tb in doc["tables"]:
        prose = prose.replace(tb, ' ')
    body = re.split(r'(?i)appendix of work', prose)[0]
    wc = len(body.split())
    ev = [f"pitch word count (appendix excluded) = {wc}"]
    hedge = re.findall(r'(?i)both options are (?:good|valid|strong)|either could|'
                       r'it depends entirely|i could see both', text)
    stance = re.findall(r'\bI (?:choose|would choose|pick|recommend|submit|argue for)\b[^.]{0,60}|'
                        r'\border\s+[AB]\s+is\s+the\s+(?:right|better|correct)\b', text, re.I)
    if wc < 250:
        return _lvl(("Not Yet", "HIGH", ev + ["under the 250-word floor — Challenge 2 requirement"]))
    if hedge and not stance:
        return _lvl(("Not Yet", "HIGH", ev + [f"hedges both ways: {hedge[:2]}"]))
    if not stance:
        return _lvl(("Not Yet", "MEDIUM", ev + [
            "no unambiguous position statement detected. THE most common failure mode is a pitch "
            "that computes correctly and concludes nothing — but detection is by cue phrase, so a "
            "guide should confirm"]))
    ev.append(f"position statement: {stance[0].strip()[:80]!r}")
    counter = re.findall(r'(?i)the (?:strongest|best) (?:argument|case) against|'
                         r'against my (?:own )?position|critics? (?:would|will) say|'
                         r'what would change my mind|i could be wrong if', text)
    if counter:
        ev.append(f"names an argument against its own position: {counter[0][:60]!r}")
        lvl, conf = "Exemplary", "LOW"
        ev.append("Exemplary on C2 requires that the counter-argument is actually ANSWERED, not "
                  "just named. LOW confidence — a guide must read it.")
    else:
        lvl, conf = "Proficient", "MEDIUM"
        ev.append("no self-counter-argument detected, so not Exemplary")
    if tier_note and tier_note[0] == "C2":
        ok = bool(re.search(tier_note[2], text))
        ev.append(f"tier gate — {tier_note[1]}: {'MET' if ok else 'NOT EVIDENCED'}")
        if not ok and lvl != "Not Yet":
            lvl, conf = "Not Yet", "LOW"
            ev.append("capped at Not Yet by the tier adjustment table; LOW confidence, cue-phrase "
                      "detection only")
    return _lvl((lvl, conf, ev))


def _c3(doc, c, checks, tier_note):
    agree = [x for x in checks if x['status'] == 'AGREES']
    check = [x for x in checks if x['status'] == 'CHECK']
    na = [x for x in checks if x['status'] == 'NOT ADDRESSED']
    ev = [f"recomputed {len(checks)} case quantities: {len(agree)} agree, {len(check)} to check, "
          f"{len(na)} not addressed"]
    for x in agree[:4]:
        ev.append(f"AGREES · {x['label']} {x['field']} = {x['value']}")
    for x in check:
        ev.append(f"CHECK THIS (advisory, not a finding of error) · {x['note']}")
    # units
    if re.search(r'(?i)\bcm3|cm³|litre|liter|\bL\b|kg|hours?\b|doses?\b', doc["text"]):
        ev.append("units are carried in the prose")
    else:
        ev.append("no units detected anywhere in the submission")
    if check:
        lvl, conf = "Not Yet", "LOW"
        ev.append("HARD LIMIT 1 — a machine BLOCKER on C3 is ADVISORY, not final. The Hero may be "
                  "right and the check may be wrong, especially where the Hero declared their own "
                  "assumption. The guide adjudicates. Do not report this to the learner as an error.")
    elif not agree:
        lvl, conf = "Not Yet", "LOW"
        ev.append("no recomputable quantity was found in the submission at all — either the work "
                  "is missing or the labels do not match the case's")
    else:
        lvl, conf = "Proficient", "MEDIUM"
        selfcheck = re.search(r"(?i)sanity check|order of magnitude|that can'?t be right|"
                              r"doesn'?t seem right|checked (?:this|it) (?:a second|another) way|"
                              r"cross[- ]check|i got it wrong|first wrong|my first (?:wrong )?attempt",
                              doc["text"])
        if selfcheck:
            lvl, conf = "Exemplary", "LOW"
            ev.append(f"visible self-check: {selfcheck.group(0)!r} — Exemplary requires that the "
                      f"check actually CAUGHT something; guide must confirm")
    if tier_note and tier_note[0] == "C3":
        ok = bool(re.search(tier_note[2], doc["text"]))
        ev.append(f"tier gate — {tier_note[1]}: {'MET' if ok else 'NOT EVIDENCED'}")
        if not ok and lvl != "Not Yet":
            lvl, conf = "Not Yet", "LOW"
            ev.append("capped at Not Yet by the tier adjustment table; cue-phrase detection only")
    return _lvl((lvl, conf, ev))


def _c4(doc, c):
    ds = c.get('deliverables', [])
    if not ds:
        return _lvl(("Not Yet", "LOW", ["case.json declares no deliverables[] — C4 cannot be scored"]))
    text = doc["text"]
    present, missing = [], []
    for d in ds:
        did = d['id']
        by_id = re.search(rf'(?<![A-Za-z0-9]){re.escape(did)}(?![0-9])', text)
        words = [w for w in re.findall(r'[a-z]{5,}', (d.get('text') or '').lower())][:8]
        hits = sum(1 for w in set(words) if w in text.lower())
        if by_id or (words and hits / len(set(words)) >= 0.5):
            present.append(did)
        else:
            missing.append(did)
    ev = [f"deliverables present: {present or '-'}", f"deliverables missing: {missing or '-'}",
          "presence detected by id and keyword overlap; whether each is 'in a form a reader can "
          "actually use' is a guide judgment"]
    if missing:
        return _lvl(("Not Yet", "MEDIUM", ev))
    return _lvl(("Proficient", "MEDIUM", ev + [
        "Exemplary requires a deliverable that goes beyond the ask IN SERVICE of the argument — "
        "not detectable mechanically"]))


def _c5(doc, case_md):
    """Originality check: does the visual's underlying data appear VERBATIM in the
    case appendices (copied) or is it derived (original)?"""
    n_vis = doc["images"] + doc["charts"]
    ev = [f"visuals detected: {doc['images']} image(s), {doc['charts']} chart(s)"]
    if n_vis == 0:
        return _lvl(("Not Yet", "MEDIUM", ev + [
            "no original visual found. If the Hero submitted a visual in a format the machine "
            "cannot read (a photo of paper, a link to a Desmos graph), this is a false negative — "
            "check the appendix link."]))
    case_nums = numbers_in(case_md)
    pool = set(doc["chart_numbers"])
    src = "chart series ranges"
    if not pool:
        for tb in doc["tables"]:
            pool |= numbers_in(tb)
        src = "tables adjacent to the visual"
    if not pool:
        ev.append("could not extract the visual's underlying data — originality UNDETERMINED")
        return _lvl(("Proficient", "LOW", ev))
    copied = {v for v in pool if any(abs(v - k) < 1e-9 for k in case_nums)}
    frac = len(copied) / len(pool)
    ev.append(f"underlying data ({src}): {len(pool)} value(s), {len(copied)} appear VERBATIM in "
              f"the case appendices = {frac:.0%} copied")
    if frac >= 0.8:
        return _lvl(("Not Yet", "MEDIUM", ev + [
            "the visual re-displays the case's own appendix data. Re-typing Appendix A into a "
            "chart is not an original visual display (RUBRIC C5)."]))
    if frac <= 0.5:
        return _lvl(("Proficient", "MEDIUM", ev + [
            f"majority of the plotted values are the Hero's own: {sorted(pool - copied)[:8]}",
            "Exemplary requires the visual to make the argument BY ITSELF — a guide must look at it"]))
    return _lvl(("Proficient", "LOW", ev + [
        "mixed: part copied, part derived. Guide should decide whether the derived part carries "
        "the argument."]))


def _c6(discussions, learner, case):
    mine = [d for d in discussions
            if d.get('learner_id') == learner and d.get('case_id') == case
            and not d.get('_example')]
    n = len(mine)
    ev = [f"{n} discussion record(s) in submissions.json for this learner and case"]
    if n == 0:
        return _lvl(("Not Yet", "LOW", ev + [
            "ZERO records is ambiguous: it means 'not recorded', not 'did not participate'. "
            "RUBRIC open question — C6 depends on data nobody currently records habitually. "
            "Treat as UNKNOWN and ask the guide, do not report a zero to the learner."]))
    engages = [d for d in mine if re.search(r'(?i)respond\w*|replied|built on|disagreed with|'
                                            r'pushed back|another hero|challenged', d.get('note') or '')]
    if n < 2:
        return _lvl(("Not Yet", "MEDIUM", ev + ["Challenge 5 requires two or more"]))
    ev.append(f"{len(engages)} record(s) note engagement with another Hero's reasoning")
    if not engages:
        return _lvl(("Not Yet", "LOW", ev + [
            "Proficient requires at least one contribution that engages another Hero's reasoning "
            "rather than restating their own; the notes do not record that. Almost certainly a "
            "recording gap rather than a learner gap — ask the guide."]))
    return _lvl(("Proficient", "MEDIUM", ev + [
        "Exemplary requires changing position under evidence and saying why, or holding under "
        "pressure and saying why — not visible in a contribution count"]))


def _c7(doc):
    """HARD LIMIT 2 — the machine may NEVER mark C7 Exemplary."""
    text = doc["text"]
    metric = re.search(r'(?i)(?:the )?metric (?:I|we) (?:optimi[sz]ed|chose|used)|'
                       r'I (?:optimi[sz]ed|measured|maximi[sz]ed) for|'
                       r'the (?:right|real) question (?:here )?is|'
                       r'what (?:we|I) (?:are|am) actually (?:choosing|deciding)', text)
    alt = re.search(r'(?i)a different metric|another metric|if (?:you|we) (?:instead )?'
                    r'(?:measured|optimi[sz]ed|counted)|under a different (?:measure|metric)|'
                    r'would give a different answer', text)
    ev = []
    if metric:
        ev.append(f"names the metric optimised: {metric.group(0)[:70]!r}")
    if alt:
        ev.append(f"acknowledges an alternative metric: {alt.group(0)[:70]!r}")
    if metric and alt:
        lvl, conf = "Proficient", "LOW"
    elif metric or alt:
        lvl, conf = "Proficient", "LOW"
        ev.append("only one half of the Proficient bar is evidenced — borderline")
    else:
        lvl, conf = "Not Yet", "LOW"
        ev.append("no evidence the Hero noticed the question was theirs to choose")
    ev.append("HARD LIMIT 2 — the machine may NEVER mark C7 Exemplary. Recognising that someone "
              "reframed the question well requires reading an argument, not scanning one. "
              "Ceiling clamped to Proficient.")
    ev.append("C7 is the criterion the machine is guessing at entirely (RUBRIC). Always LOW.")
    assert lvl != "Exemplary"
    return _lvl((lvl, conf, ev))


# ------------------------------------------------------------------ score ---
def score(project, learner, case, files, note=None):
    cdir = os.path.join(project, 'cases', case)
    c = json.load(open(os.path.join(cdir, 'case.json'), encoding='utf-8'))
    case_md = open(os.path.join(cdir, 'case.md'), encoding='utf-8').read()
    subs = json.load(open(os.path.join(project, 'readmodel', 'submissions.json'), encoding='utf-8'))

    doc = ingest(files)
    tier = c.get('tier')
    tier_note = TIER_GATE.get(tier)
    table = value_table(case_md, c)
    checks = recompute(doc["text"], table)

    tracker_sub = None
    for s in subs.get('submissions', []):
        if s.get('learner_id') == learner and s.get('case_id') == case and not s.get('_example'):
            tracker_sub = s.get('artifacts')

    glossary_urls = case_glossary_urls(cdir)

    v = {
        "C1": _c1(doc, c, tracker_sub, glossary_urls),
        "C2": _c2(doc, c, tier_note),
        "C3": _c3(doc, c, checks, tier_note),
        "C4": _c4(doc, c),
        "C5": _c5(doc, case_md),
        "C6": _c6(subs.get('discussions', []), learner, case),
        "C7": _c7(doc),
    }
    assert v["C7"]["level"] != "Exemplary", "HARD LIMIT 2 violated"

    # MCS-31 enforcement. A criterion the machine gets wrong >30% of the time
    # (guide_confirm.py --audit --write) is demoted to GUIDE_ONLY: the machine
    # does not guess, it refuses. A demoted criterion fails the mastery gate
    # until a guide supplies a real level via guide_confirm.py --override, and
    # contributes 0 to the growth score in the interim (never a silent pass).
    demoted = {}
    dc_path = os.path.join(project, 'readmodel', 'demoted_criteria.json')
    if os.path.exists(dc_path):
        demoted = json.load(open(dc_path, encoding='utf-8')).get('demoted', {})
    for k, why in demoted.items():
        if k in v:
            v[k] = {"level": None, "confidence": "GUIDE_ONLY",
                    "evidence": [f"DEMOTED per MCS-31: machine wrong {why['rate']:.0%} of "
                                 f"{why['n']} confirmed records (demoted {why['demoted_on']}). "
                                 f"The machine does not score this criterion. A guide must "
                                 f"set it directly via guide_confirm.py --override."]}

    def _score_of(level):
        return LEVELS.get(level, 0) if level is not None else 0

    badge = all(v[k]["level"] is not None and LEVELS[v[k]["level"]] >= 70 for k in WEIGHTS)
    growth = round(sum(WEIGHTS[k] * _score_of(v[k]["level"]) for k in WEIGHTS) / 100.0, 1)
    lows = [k for k in WEIGHTS if v[k]["confidence"] in ("LOW", "GUIDE_ONLY")]

    return {
        "schema": "submission-score/1.0",
        "learner_id": learner, "case_id": case, "tier": tier,
        "scored_on": datetime.date.today().isoformat(),
        "status": "machine-first-pass",          # HARD LIMIT 3
        "files_read": doc["read"], "files_unreadable": doc["unreadable"],
        "criteria": v,
        "mastery_gate": {"badge_eligible": badge,
                         "not_yet": [k for k in WEIGHTS if v[k]["level"] == "Not Yet"],
                         "guide_only": [k for k in WEIGHTS if v[k]["level"] is None]},
        "growth_score": growth,
        "surface_first": lows,
        "recompute": checks,
        "tier_gate": ({"tier": tier, "criterion": tier_note[0], "requirement": tier_note[1]}
                      if tier_note else {"tier": tier, "requirement": None}),
        "note": note,
    }


def render(r):
    L = []
    A = L.append
    A(f"# SCORE — {r['learner_id']} · {r['case_id']}")
    A("")
    A(f"**Status:** `{r['status']}` · **Scored:** {r['scored_on']} · **Tier:** {r['tier']}")
    A("")
    A("> **INTERIM.** This is a machine first pass, not a verdict. It does not count toward a "
      "badge and renders as INTERIM on the dashboard until a guide confirms it. "
      "Machine drafts, human promotes.")
    A("")
    A("## The two outputs — never averaged")
    A("")
    A("| | |")
    A("|---|---|")
    _blockers = r['mastery_gate']['not_yet'] + [f"{k} (guide-only, demoted)" for k in r['mastery_gate'].get('guide_only', [])]
    A(f"| **Mastery gate** — badge eligible? | **{'YES' if r['mastery_gate']['badge_eligible'] else 'NO'}**"
      f"{' — blocked on ' + ', '.join(_blockers) if _blockers else ''} |")
    A(f"| **Growth score** (0–100, guide-only, never a grade) | **{r['growth_score']}** |")
    A("")
    A("All-Proficient scores 70. The score tracks trajectory; the gate tracks readiness. "
      "If a learner asks what their score is, the answer is that it is a tool for their guide, "
      "and the number that matters is which criteria are still Not Yet.")
    A("")
    if r['surface_first']:
        A(f"## Surface to the guide first — LOW confidence on {', '.join(r['surface_first'])}")
        A("")
    A("## The seven criteria")
    A("")
    A("| | Criterion | Weight | Verdict | Confidence |")
    A("|---|---|---|---|---|")
    for k in ["C1", "C2", "C3", "C4", "C5", "C6", "C7"]:
        A(f"| {k} | {NAMES[k]} | {WEIGHTS[k]} | **{r['criteria'][k]['level'] or 'GUIDE ONLY'}** | "
          f"{r['criteria'][k]['confidence']} |")
    A("")
    for k in ["C1", "C2", "C3", "C4", "C5", "C6", "C7"]:
        cr = r['criteria'][k]
        A(f"### {k} · {NAMES[k]} — {cr['level'] or 'GUIDE ONLY (demoted, see MCS-31)'} ({cr['confidence']} confidence)")
        for e in cr['evidence']:
            A(f"- {e}")
        A("")
    A("## Recomputation of the Hero's numbers")
    A("")
    A("> A machine BLOCKER on C3 is **advisory, not final**. Where the machine disagrees, the "
      "Hero may be right and the check may be wrong — especially where the Hero declared their "
      "own assumption. Presented as *check this*, never as *wrong*.")
    A("")
    A("| quantity | machine | Hero | status |")
    A("|---|---|---|---|")
    for x in r['recompute']:
        A(f"| {x['label']} · {x['field']} | {x['value']} | {x.get('hero', '—')} | {x['status']} |")
    A("")
    tg = r['tier_gate']
    A(f"## Tier adjustment ({tg['tier']})")
    A("")
    A(f"- Requirement: {tg['requirement'] or '— (T1 has none)'}")
    if tg.get('criterion'):
        A(f"- Gates: {tg['criterion']}")
    A("- The tier note changes what Proficient *requires*; it does **not** scale the score. "
      "A Proficient T1 and a Proficient T4 both score 70.")
    A("")
    if r['files_unreadable']:
        A("## NOT READ")
        for u in r['files_unreadable']:
            A(f"- {u}")
        A("")
        A("Unreadable is not empty. Nothing above is scored against these.")
        A("")
    A("## Guide confirm")
    A("")
    A("```bash")
    A(f"python3 skills/submission-score/scripts/guide_confirm.py \\")
    A(f"    --learner {r['learner_id']} --case {r['case_id']} --by \"<your name>\" \\")
    A(f"    --override C3=Proficient --reason \"Hero declared their own assumption; machine "
      f"recomputed against the case default\"")
    A("```")
    A("")
    A("Every override is logged to `readmodel/score_overrides.json`. **Overrides are the training "
      "signal** — a criterion the machine gets wrong more than ~30% of the time is demoted to "
      "guide-only.")
    return "\n".join(L)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--learner', required=True)
    ap.add_argument('--case', required=True)
    ap.add_argument('--file', action='append', required=True)
    ap.add_argument('--project', default=os.getcwd())
    ap.add_argument('--note', default=None)
    ap.add_argument('--write', action='store_true',
                    help='write SCORE.md and append to readmodel/submissions.json scores[]')
    a = ap.parse_args()

    r = score(a.project, a.learner, a.case, a.file, a.note)
    md = render(r)
    print(md)

    if a.write:
        d = os.path.join(a.project, 'student_work', a.learner, a.case)
        os.makedirs(d, exist_ok=True)
        open(os.path.join(d, 'SCORE.md'), 'w', encoding='utf-8').write(md + "\n")
        snapshot(a.project, ['submissions.json'])
        sp = os.path.join(a.project, 'readmodel', 'submissions.json')
        subs = json.load(open(sp, encoding='utf-8'))
        subs.setdefault('scores', [])
        rid = f"{a.learner}__{a.case}__{r['scored_on']}"
        rec = {"id": rid, "learner_id": a.learner, "case_id": a.case,
               "scored_on": r['scored_on'], "status": r['status'],
               "criteria": {k: {"level": v['level'], "confidence": v['confidence']}
                            for k, v in r['criteria'].items()},
               "badge_eligible": r['mastery_gate']['badge_eligible'],
               "growth_score": r['growth_score'],
               "surface_first": r['surface_first'],
               "confirmed_by": None, "confirmed_on": None, "overrides": [],
               "report": f"student_work/{a.learner}/{a.case}/SCORE.md"}
        subs['scores'] = [x for x in subs['scores'] if x.get('id') != rid] + [rec]
        json.dump(subs, open(sp, 'w', encoding='utf-8'), indent=2)
        print(f"\nWROTE {d}/SCORE.md and readmodel/submissions.json scores[{rid}]")


if __name__ == '__main__':
    main()
