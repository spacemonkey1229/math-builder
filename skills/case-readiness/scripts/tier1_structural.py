#!/usr/bin/env python3
"""
TIER 1 — STRUCTURAL GATE  (machine, blocking)

Nineteen mechanical checks against SPEC.md.  Supersedes the inline script that
used to live in skills/case-qa/SKILL.md, with three defect fixes and six new
continuity checks.

    python3 skills/case-readiness/scripts/tier1_structural.py --case cold-chain
    python3 skills/case-readiness/scripts/tier1_structural.py --case cold-chain --json out.json

Exit code 0 = TIER1_PASS (zero BLOCKERs).  Exit code 1 = BLOCKED.
This script NEVER writes to the case.  It only reports.

CHECK REGISTER  (name them by id in the QA report)
  T1-01  structural sections present                       SPEC §1
  T1-02  Overview length + no load-bearing numerals        SPEC §1.2
  T1-03  Dilemma is binary, 'You are <X>.'                 SPEC §1.4
  T1-04  Challenges 1/2/4/5 verbatim boilerplate           SPEC §1.5   [FIXED: numeral-form tolerant]
  T1-05  Challenge numeral-form drift                      SPEC §1.5   [NEW: reported separately]
  T1-06  appendices lettered, ordered, no orphans          SPEC §1.6   [FIXED: no "how to" blanket exemption]
  T1-07  final appendix lines are verbatim spine labels    SPEC §1.7   [FIXED: reads the dot-separated Vocabulary line]
  T1-08  every number cited or declared an assumption      SPEC §5
  T1-09  balance test                                      SPEC §2
  T1-10  standards binding / case.json schema              SPEC §4
  T1-11  learning-goal mutual reachability (MCS-10)        SPEC §4
  T1-12  tier vs numeric load                              SPEC §3
  T1-13  facilitator guide + anchor ladder + sensitivity   SPEC §7, MCS-14/15/16
  T1-14  CONTINUITY — FACILITATOR.md agrees with case.md   MCS-23      [NEW]
  T1-15  CONTINUITY — published artifacts agree/are fresh  MCS-23      [NEW]
  T1-16  QA hash currency                                  MCS-11
  T1-17  cohort-size bands declared                        MCS-25      [NEW]
  T1-18  interactive HTML present, fresh, leak-free         MCS-37      [NEW]
  T1-19  hover+hyperlink glossary, verified domains         MCS-39      [NEW]
"""
import argparse, json, os, re, sys, hashlib, unicodedata

# ---------------------------------------------------------------- helpers ----
DASH = r'[—–\-:]'          # em dash, en dash, hyphen, colon
MIDDOT = '·'                    # ·

def norm(s: str) -> str:
    s = (s.replace('“', '"').replace('”', '"')
           .replace('‘', "'").replace('’', "'"))
    return re.sub(r'\s+', ' ', s).strip()

NUM_RE = re.compile(r'(?<![\w.])(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\d+)(?![\w])')

def numbers_in(text: str):
    """Every numeric literal in a blob of text, as floats."""
    out = set()
    for m in NUM_RE.finditer(text):
        try:
            out.add(float(m.group(1).replace(',', '')))
        except ValueError:
            pass
    return out

def close(a, b, rel=1e-4):
    if a == b:
        return True
    scale = max(abs(a), abs(b), 1.0)
    return abs(a - b) / scale < rel

def in_universe(v, universe):
    return any(close(v, u) for u in universe)


class Report:
    def __init__(self):
        self.B, self.G, self.P = [], [], []
        self.checks = {}

    def blocker(self, cid, msg):
        self.B.append(f"[{cid}] {msg}"); self.checks.setdefault(cid, 'FAIL'); self.checks[cid] = 'FAIL'

    def gap(self, cid, msg):
        self.G.append(f"[{cid}] {msg}")
        if self.checks.get(cid) != 'FAIL':
            self.checks[cid] = 'GAP'

    def polish(self, cid, msg):
        self.P.append(f"[{cid}] {msg}")
        self.checks.setdefault(cid, 'PASS')

    def ok(self, cid):
        self.checks.setdefault(cid, 'PASS')


# ------------------------------------------------------------------ main ----
def run(project, cid):
    R = Report()
    cdir = os.path.join(project, 'cases', cid)
    md_path = os.path.join(cdir, 'case.md')
    md = open(md_path, encoding='utf-8').read()
    c = json.load(open(os.path.join(cdir, 'case.json'), encoding='utf-8'))
    sp = json.load(open(os.path.join(project, 'readmodel', 'standards_spine.json'), encoding='utf-8'))
    bp = open(os.path.join(project, 'doctrine', 'CHALLENGE_BOILERPLATE.md'), encoding='utf-8').read()

    labels, by_id = {}, {}
    for dom in sp['domains']:
        for g in dom['groups']:
            for s in g['skills']:
                by_id[s['id']] = s
                labels.setdefault(s['label'], []).append(s['id'])

    # ---- T1-01 structural sections -----------------------------------------
    for name, pat in [("Title block", r'Case Study:\s*\S'),
                      ("MS banner",   r'MS\s*\|\s*Math\s*\|\s*Case Study'),
                      ("Overview",    r'(?im)^#+\s*Overview'),
                      ("Background",  r'(?im)^#+\s*Background'),
                      ("Dilemma",     r'(?im)^#+\s*(The )?Dilemma'),
                      ("Challenges",  r'(?im)^#+\s*Challenges'),
                      ("Appendices",  r'(?im)^#+\s*Appendix\s+A')]:
        if not re.search(pat, md):
            R.blocker('T1-01', f"SPEC §1 — missing section: {name}")
    R.ok('T1-01')

    # ---- T1-02 Overview -----------------------------------------------------
    # A numeral in the Overview is BLOCKER-for-confirmation, per SPEC §1.2 — but
    # two forms are mechanically classifiable as texture and are auto-downgraded:
    # a clock time (4:40) and a 4-digit year.  Anything else must be cleared by
    # hand and the clearance recorded in case.json.overview_numerals_cleared[].
    ov = re.search(r'(?is)^#+\s*Overview\s*(.*?)(?=^#+\s)', md, re.M)
    if ov:
        ovt = ov.group(1)
        words = len(ovt.split())
        if not 150 <= words <= 300:
            R.gap('T1-02', f"SPEC §1.2 — Overview is {words} words, spec says 150-300")
        cleared = {str(x.get('numeral') if isinstance(x, dict) else x)
                   for x in (c.get('overview_numerals_cleared') or [])}
        texture, suspect = [], []
        for m in NUM_RE.finditer(ovt):
            tok = m.group(1)
            lo, hi = max(0, m.start() - 3), min(len(ovt), m.end() + 3)
            ctx = ovt[lo:hi]
            is_clock = bool(re.search(r'\d{1,2}:\d{2}', ctx))
            is_year = bool(re.fullmatch(r'(1[89]|20)\d\d', tok))
            (texture if (is_clock or is_year or tok in cleared) else suspect).append(tok)
        if suspect:
            R.blocker('T1-02', f"SPEC §1.2 — Overview contains numerals {suspect[:6]} that are neither "
                               f"clock times nor years. Read them. If they are texture, record the "
                               f"clearance in case.json.overview_numerals_cleared[]; if a student could "
                               f"use one in a calculation it stays a BLOCKER.")
        if texture:
            R.polish('T1-02', f"Overview numerals auto-classified as texture (clock/year/cleared): {texture}")
    R.ok('T1-02')

    # ---- T1-03 Dilemma ------------------------------------------------------
    dl = re.search(r'(?is)^#+\s*(?:The )?Dilemma\s*(.*?)(?=^#+\s)', md, re.M)
    if dl:
        d = dl.group(1)
        if not re.search(r'(?m)^\s*(?:\*\*)?A\.\)', d) or not re.search(r'(?m)^\s*(?:\*\*)?B\.\)', d):
            R.blocker('T1-03', "SPEC §1.4 — Dilemma must use the A.) / B.) form")
        if re.search(r'(?m)^\s*(?:\*\*)?C\.\)', d):
            R.blocker('T1-03', "SPEC §1.4 — Dilemma is not binary (found option C)")
        if not re.search(r'(?i)you are\s+\S', d):
            R.blocker('T1-03', "SPEC §1.4 — Dilemma must open 'You are <Protagonist>.'")
        bgm = re.search(r'(?is)^#+\s*Background\s*(.*?)(?=^#+\s)', md, re.M)
        for k in ("A", "B"):
            opt = (c.get('options') or {}).get(k) or ""
            if opt and bgm:
                key = [w for w in norm(opt).split()[:6] if len(w) > 4]
                if key and not any(w.lower() in norm(bgm.group(1)).lower() for w in key):
                    R.gap('T1-03', f"SPEC §1.3 — Option {k} may not be named in Background (fuzzy match; "
                                   f"read it and clear or keep deliberately)")
    R.ok('T1-03')

    # ---- T1-04 / T1-05 Challenge boilerplate --------------------------------
    # FIX #1: the old regex hard-required `N.)`.  A case written with `1.` or
    # `**1.**` slipped the whole diff and a rewritten Challenge 2 passed silently
    # (that is real: cold-chain QA_REPORT B3).  We now accept `N.)`, `N.`, `N)`,
    # with or without bold, AND report the form used as a separate finding so the
    # drift is still visible instead of being normalised away.
    CH_HDR = re.compile(r'(?m)^[ \t]*(?P<b1>\*\*)?(?P<n>[1-5])(?P<sep>\.\)|\.|\))(?P<b2>\*\*)?[ \t]+')
    RULE_TAIL = re.compile(r'(?:\s*(?:-{3,}|\*{3,}|_{3,}))+\s*$')
    SUBMARK = re.compile(r'(?m)^[ \t]*[-*+][ \t]+(?=[a-z]\.\))')     # '- a.)' -> 'a.)'

    def strip_sub_markers(t):
        return SUBMARK.sub('', t)

    canon = re.search(r'## BEGIN CANONICAL BLOCK(.*?)## END CANONICAL BLOCK', bp, re.S).group(1)
    canon = re.search(r'```(.*?)```', canon, re.S).group(1)
    LOCKED_NS = (1, 2, 4, 5)          # SPEC §1.5 — only these four are verbatim
    locked, canon_form = {}, {}
    for m in list(CH_HDR.finditer(canon)):
        n = int(m.group('n'))
        canon_form[n] = (m.group('b1') or '') + m.group('n') + m.group('sep') + (m.group('b2') or '')
        nxt = CH_HDR.search(canon, m.end())
        if n in LOCKED_NS:
            locked[n] = norm(RULE_TAIL.sub('', canon[m.end(): nxt.start() if nxt else len(canon)]))

    ch = re.search(r'(?is)^#+\s*Challenges\s*(.*?)(?=^#+\s*Appendix|\Z)', md, re.M)
    if not ch:
        R.blocker('T1-04', "SPEC §1.5 — Challenges section not found")
    else:
        have = ch.group(1)
        hdrs = list(CH_HDR.finditer(have))
        seen = {}
        for i, m in enumerate(hdrs):
            n = int(m.group('n'))
            if n in seen:
                continue
            nxt = hdrs[i + 1].start() if i + 1 < len(hdrs) else len(have)
            raw = RULE_TAIL.sub('', have[m.end():nxt])
            seen[n] = ((m.group('b1') or '') + m.group('n') + m.group('sep') + (m.group('b2') or ''),
                       norm(raw), norm(strip_sub_markers(raw)))
        for n, text in locked.items():
            if n not in seen:
                R.blocker('T1-04', f"SPEC §1.5 — Challenge {n} missing")
                continue
            form, got_raw, got = seen[n]
            if form != canon_form[n]:
                R.gap('T1-05', f"SPEC §1.5 — Challenge {n} numeral form is {form!r}, boilerplate locks "
                               f"{canon_form[n]!r}. The text was still diffed (the legacy check would "
                               f"have skipped the whole section). Restore the locked form.")
            if got_raw != got and got == text:
                R.gap('T1-05', f"SPEC §1.5 — Challenge {n} sub-lines are rendered as markdown bullets "
                               f"('- a.)'); the boilerplate locks a bare indented 'a.)'. Text is otherwise "
                               f"verbatim. Rendering drift, not a rewrite.")
            if got != text:
                i0 = next((k for k in range(min(len(text), len(got))) if text[k] != got[k]),
                          min(len(text), len(got)))
                a, b = max(0, i0 - 40), i0 + 70
                R.blocker('T1-04', f"SPEC §1.5 — Challenge {n} is NOT verbatim boilerplate "
                                   f"(diverges at char {i0}).\n"
                                   f"      expected: ...{text[a:b]}...\n      found:    ...{got[a:b]}...")
        if 3 not in seen:
            R.blocker('T1-04', "SPEC §1.5 — Challenge 3 (case-specific deliverables) missing")
        elif seen[3][0] != canon_form[3]:
            R.gap('T1-05', f"SPEC §1.5 — Challenge 3 heading numeral is {seen[3][0]!r}, locked form is "
                           f"{canon_form[3]!r} (contents are free, the numeral is not)")
    R.ok('T1-04'); R.ok('T1-05')

    # ---- T1-06 appendices ---------------------------------------------------
    # FIX #2: the old check did `if re.search('how to', title): continue`, which
    # exempted ANY appendix whose title mentions "how to" from the orphan graph.
    # Only the single mandatory final standards appendix (SPEC §1.7) gets relief,
    # and even that is a GAP rather than silence.
    apps = re.findall(r'(?im)^#+\s*Appendix\s+([A-Z])\s*' + DASH + r'?\s*(.+)$', md)
    letters = [a for a, _ in apps]
    if letters != sorted(letters):
        R.gap('T1-06', f"Appendices out of order: {letters}")
    if len(letters) != len(set(letters)):
        R.blocker('T1-06', f"Duplicate appendix letters: {letters}")
    FINAL_TITLE = re.compile(r'(?i)how\s*to.{0,4}\s*(math\s+)?questions?\s+to\s+research')
    final_letter = None
    for L, title in apps:
        if FINAL_TITLE.search(title):
            final_letter = L                     # last one wins; §1.7 says it is final
    body_wo_apps = re.split(r'(?im)^#+\s*Appendix\s+A\b', md)[0]
    for L, title in apps:
        referenced = bool(re.search(rf'(?i)APPENDI(?:X|CES)[^.\n]*\b{L}\b', body_wo_apps)) or \
                     bool(re.search(rf'(?i)\*\*APPENDIX\s+{L}\*\*', md))
        if referenced:
            continue
        if L == final_letter:
            R.gap('T1-06', f"SPEC §1.6/§1.7 — final standards Appendix {L} ({title.strip()}) is not "
                           f"referenced by letter anywhere in Background or Challenges. Challenge 1 points "
                           f"at it by construction, so this is a GAP, not a BLOCKER — but name it explicitly.")
        else:
            R.blocker('T1-06', f"SPEC §1.6 — Appendix {L} ({title.strip()}) is an ORPHAN: never "
                               f"referenced from Background or Challenges")
    R.ok('T1-06')

    # ---- T1-07 final appendix verbatim spine labels -------------------------
    # FIX #3: the old check only read lines starting '- ' or '* '.  Cold Chain's
    # Vocabulary block is a single middot-separated run — 14 labels the old check
    # skipped entirely, so MCS-05 was only ever enforced on the "How to..." half.
    fin = re.search(r'(?is)^#+\s*Appendix\s+[A-Z]\s*' + DASH + r'?\s*.{0,12}How To.*?$(.*)\Z', md, re.M)
    if not fin:
        R.blocker('T1-07', 'SPEC §1.7 — mandatory final appendix ("How To..." Math Questions To Research) missing')
    else:
        body = fin.group(1)
        lines, bullet_n, vocab_n = [], 0, 0
        for raw in body.splitlines():
            t = norm(raw)
            if not t:
                continue
            if t.startswith(('- ', '* ', '+ ')):
                lines.append(t[2:].strip().strip('*').strip()); bullet_n += 1
            elif t.startswith('#') or (t.startswith('**') and t.endswith('**') and MIDDOT not in t):
                continue                       # a sub-heading like **Vocabulary**
            elif MIDDOT in t or '•' in t:
                for part in re.split(r'[' + MIDDOT + r'•]', t):
                    p = part.strip().strip('*').strip()
                    if p:
                        lines.append(p); vocab_n += 1
        if not lines:
            R.blocker('T1-07', "SPEC §1.7 — final appendix has no skill lines")
        bad = [l for l in lines if l not in labels]
        if bad:
            R.blocker('T1-07', f"SPEC §1.7 — final-appendix lines are NOT verbatim spine labels: {bad}")
        claimed = {by_id[i]['label'] for i in c.get('standards', []) if i in by_id}
        miss = sorted(claimed - set(lines))
        extra = sorted(set(lines) - claimed - set(bad))
        if miss:
            R.gap('T1-07', f"standards[] claims skills absent from the final appendix: {miss}")
        if extra:
            R.gap('T1-07', f"final appendix lists skills absent from standards[]: {extra}")
        if not re.search(r'(?i)vocabulary', body) or not re.search(r'how to', body, re.I):
            R.gap('T1-07', 'SPEC §1.7 — final appendix should be a two-part list '
                           '(Vocabulary terms / "How to..." questions)')
        R.polish('T1-07', f"parsed {len(lines)} labels ({bullet_n} bulleted, {vocab_n} from "
                          f"separator-run lines) — the separator run is the half the legacy check skipped")
    R.ok('T1-07')

    # ---- T1-08 every number cited or declared -------------------------------
    left = re.findall(r'\{\{NEEDS_SOURCE:[^}]*\}\}', md)
    if left:
        R.blocker('T1-08', f"SPEC §5 — {len(left)} unresolved {{{{NEEDS_SOURCE}}}} placeholder(s) remain")
    ma = re.search(r'(?is)^#+\s*Appendix\s+[A-Z]\s*' + DASH + r'?\s*Modeling Assumptions\s*(.*?)(?=^#+\s|\Z)', md, re.M) \
         or re.search(r'(?is)^#+\s*Modeling Assumptions\s*(.*?)(?=^#+\s|\Z)', md, re.M)
    ma_txt = ma.group(1) if ma else ""
    src_ids = {s.get('id') for s in c.get('sources', [])}
    marks = set(re.findall(r'\[(S\d+)\]', md))
    if marks - src_ids:
        R.blocker('T1-08', f"SPEC §5 — citation markers with no sources[] row: {sorted(marks - src_ids)}")
    for s in c.get('sources', []):
        missk = [k for k in ("claim", "value", "org", "title", "url", "accessed", "confidence") if not s.get(k)]
        if missk:
            R.blocker('T1-08', f"SPEC §5 — source {s.get('id')} incomplete, missing {missk}")
        if s.get('confidence') not in ("HIGH", "MEDIUM", "LOW"):
            R.blocker('T1-08', f"source {s.get('id')} has no valid confidence grade")
    if [s for s in c.get('sources', []) if s.get('confidence') == 'LOW']:
        R.gap('T1-08', "LOW-confidence sources present — confirm none carries a constraint or a "
                       "dilemma-flipping quantity")
    declared_assumptions = c.get('modeling_assumptions') or c.get('assumptions')
    if declared_assumptions and not ma_txt.strip():
        R.blocker('T1-08', "SPEC §5 — case.json declares assumptions but case.md has no visible "
                           "'Modeling Assumptions' block")
    # A legacy conversion whose own methodology note claims its numbers are "verbatim" from a
    # real source or "publicly verifiable" is claiming something externally checkable -- e.g. a
    # real clinical trial's mortality table, a real still-active athlete's career stats. Tier 1
    # cannot verify the claim itself (that needs Tier 2 P6 / a web-connected fact-check), but it
    # CAN verify the claim is at least backed by a sources[] entry, the same way T1-08 already
    # requires a citation marker to resolve to a sources[] row. An empty sources[] here means
    # nothing was ever checked against reality -- exactly how a fabricated clinical trial table
    # and a stale athlete season both passed Tier 1 clean before this check existed (2026-08-04).
    REAL_CLAIM_RE = re.compile(r'verbatim|publicly verifiable', re.I)
    legacy_note = (c.get('legacy_conversion_note') or '')
    if c.get('origin') == 'legacy' and REAL_CLAIM_RE.search(legacy_note):
        real_sources = [s for s in c.get('sources', []) if s.get('confidence') in ('HIGH', 'MEDIUM')]
        if not real_sources:
            R.blocker('T1-08', "legacy_conversion_note claims data is 'verbatim'/'publicly verifiable' "
                               "from a real source, but sources[] has no HIGH/MEDIUM-confidence entry to "
                               "back it. A claim of real, externally-checkable data with nothing checking "
                               "it is exactly the gap that let fabricated/stale real-world data reach "
                               "TIER1_PASS before -- add a sources[] entry (or run the P6 data-decay probe "
                               "and record it) before this can pass.")
    R.ok('T1-08')

    # ---- T1-09 balance test -------------------------------------------------
    bal = c.get('balance') or c.get('balance_test') or {}
    arg_a = bal.get('argument_a') or bal.get('argument_A') or ''
    arg_b = bal.get('argument_b') or bal.get('argument_B') or ''
    if not arg_a.strip():
        R.blocker('T1-09', "SPEC §2 — strongest quantitative argument for A not written")
    if not arg_b.strip():
        R.blocker('T1-09', "SPEC §2 — strongest quantitative argument for B not written")
    mp = c.get('margin_pct', bal.get('margin_pct'))
    if mp is None:
        R.blocker('T1-09', "SPEC §2 — margin_pct not recorded")
    elif mp > 20:
        R.blocker('T1-09', f"SPEC §2 — UNBALANCED: margin is {mp}% on '{c.get('primary_metric')}' "
                           f"(cap ~20%). Retune a constraint. Do NOT fix this by hiding data.")
    elif mp > 15:
        R.gap('T1-09', f"margin {mp}% is inside the cap but wide; the discussion will be one-sided")
    if not (c.get('primary_metric') or '').strip():
        R.blocker('T1-09', "SPEC §2/§4 — primary_metric not stated")
    R.ok('T1-09')

    # ---- T1-10 standards binding -------------------------------------------
    for k in ("id", "title", "tier", "band", "primary_metric", "options", "margin_pct",
              "standards", "new_coverage", "sources"):
        if k not in c:
            R.blocker('T1-10', f"SPEC §4 — case.json missing key: {k}")
    for s in c.get('standards', []):
        if s not in by_id:
            R.blocker('T1-10', f"SPEC §4 — standards[] id not in spine: {s}")
    for s in c.get('new_coverage', []):
        if s not in c.get('standards', []):
            R.blocker('T1-10', f"new_coverage[] not a subset of standards[]: {s}")
        else:
            skill = by_id.get(s, {})
            n = skill.get('n', 0)
            covered_by = skill.get('covered_by', [])
            title = c.get('title')
            # n==0 (never covered) is always a valid new claim. n==1 is valid ONLY
            # if the sole contributor is this case's own title -- i.e. this case's
            # own earlier promotion already bumped 0->1 and we're re-checking after
            # the fact (MCS: documented false-positive-after-own-promotion pattern).
            # n==1 from a DIFFERENT case is not that: it means someone else covered
            # this skill before this case ever claimed it, which is exactly the bug
            # that let Nuclear Waste's original 'geometry-prism' claim (n=1, but the
            # 1 was Cold Chain, not itself) pass this check even though the claim
            # was never actually new. Never accept n>=2 regardless of who's in it.
            valid = (n == 0) or (n == 1 and covered_by == [title])
            if not valid:
                R.blocker('T1-10', f"new_coverage[] claims '{s}' but the spine shows n={n} "
                                   f"(covered_by={covered_by})")
    if not c.get('new_coverage'):
        # Legacy-conversion carve-out (added 2026-08-13, project-owner approved):
        # a case with a legacy_conversion_note is being FORMALIZED, not freshly
        # authored -- its title was already tagged into standards_spine.json's
        # covered_by[] by the pre-pipeline corpus scan, so every skill its
        # deliverables genuinely force can legitimately already have n>0. SPEC §8's
        # "a case that closes no gap has no reason to exist" is aimed at a fresh
        # case with no pedagogical purpose; it is the wrong test for a corpus case
        # that already has one. This does NOT relax T1-11 (mutual reachability) or
        # the anti-padding rule above (n>=2 claims are still always a BLOCKER) --
        # a legacy case still cannot claim a skill its deliverables don't force,
        # and still cannot claim a skill covered by another case's title. It only
        # downgrades the single "closes no NEW gap" finding from BLOCKER to a
        # recorded GAP, for legacy cases specifically.
        if c.get('legacy_conversion_note'):
            R.gap('T1-10', "SPEC §8 — new_coverage[] is empty; the case closes no NEW gap "
                            "(legacy-conversion carve-out: this title is already credited "
                            "in standards_spine.json's covered_by[] from the pre-pipeline "
                            "corpus scan, so no uncovered skill remains for it to honestly "
                            "claim -- downgraded from BLOCKER per 2026-08-13 doctrine "
                            "amendment, project-owner approved, not a per-case waiver)")
        else:
            R.blocker('T1-10', "SPEC §8 — new_coverage[] is empty; the case closes no gap")
    R.ok('T1-10')

    # ---- T1-11 learning-goal MUTUAL reachability (MCS-10) -------------------
    delivs = c.get('deliverables', [])
    mapped = {sk for d in delivs for sk in d.get('skills', [])}
    unmapped = [d['id'] for d in delivs if not d.get('skills')]
    unreach = sorted(set(c.get('standards', [])) - mapped)
    phantom = sorted(mapped - set(c.get('standards', [])))     # <- the other direction
    if not delivs:
        R.blocker('T1-11', "MCS-10 — case.json has no deliverables[]; the reachability check cannot run, "
                           "which means every claimed standard is unreachable by construction")
    if unmapped:
        R.blocker('T1-11', f"LEARNING GOAL — deliverables mapped to no skill: {unmapped}")
    if unreach:
        hits_new = set(unreach) & set(c.get('new_coverage', []))
        msg = f"LEARNING GOAL — skills no deliverable forces: {unreach}"
        if hits_new:
            R.blocker('T1-11', msg + " (includes new_coverage — this would falsely close a gap)")
        else:
            R.gap('T1-11', msg)
    if phantom:
        R.blocker('T1-11', f"MCS-10 (reverse direction) — deliverables claim skills that are not in "
                           f"standards[]: {phantom}. The deliverable forces work the case never claims "
                           f"credit for, so the coverage matrix understates the case.")
    R.ok('T1-11')

    # ---- T1-12 tier vs numeric load ----------------------------------------
    tier = c.get('tier')
    load = {"sci": bool(re.search(r'\d\s*[x×]\s*10\^?\d|e[+-]\d\d|scientific notation', md, re.I)),
            "prob": bool(re.search(r'(?i)probabilit|likelihood|odds', md)),
            "assume": bool(ma_txt.strip())}
    nvars = len(delivs)
    if tier == "T4" and not load["assume"]:
        R.blocker('T1-12', "SPEC §3 — T4 requires at least one quantity the student must assume, stated "
                           "and justified; no Modeling Assumptions block found")
    if tier == "T1" and (load["sci"] or load["prob"] or nvars > 3):
        R.gap('T1-12', "SPEC §3 — declared T1 but the numeric load looks T2+; re-tier")
    if tier in ("T3", "T4") and nvars <= 2 and not (load["sci"] or load["prob"]):
        R.gap('T1-12', f"SPEC §3 — declared {tier} but only {nvars} deliverables and no advanced "
                       f"operations; may be T2")
    if tier not in ("T1", "T2", "T3", "T4"):
        R.blocker('T1-12', "SPEC §3 — tier must be T1-T4")
    R.ok('T1-12')

    # ---- T1-13 facilitator artifacts ---------------------------------------
    fac_path = os.path.join(cdir, 'FACILITATOR.md')
    fac = ''
    if not os.path.exists(fac_path):
        R.blocker('T1-13', "SPEC §7 — FACILITATOR.md missing")
    else:
        fac = open(fac_path, encoding='utf-8').read()
        qs = len(re.findall(r'(?m)^\s*(?:\d+[.)]|[-*>])\s.*\?', fac))
        if qs < 5:
            R.gap('T1-13', f"facilitator guide has {qs} Socratic questions, spec asks for 5")
        if not re.search(r'(?i)trap', fac):
            R.gap('T1-13', "facilitator guide does not name 'the trap'")
        # MCS-14 anchor ladder, now mechanical
        ANCHORS = ["Intro", "Launch", "Anchor 1", "Anchor 2", "Anchor 3", "Anchor 4",
                   "Data Deep Dive", "Twist", "Close", "Lessons Learned", "Standards Touched On"]
        missing = [a for a in ANCHORS if not re.search(r'(?im)^#+.*' + re.escape(a), fac)]
        if missing:
            R.blocker('T1-13', f"MCS-14 — anchor ladder incomplete, missing headings: {missing}")
        # order: Anchors 2 and 3 must precede the Data Deep Dive
        pos = {a: (re.search(r'(?im)^#+.*' + re.escape(a), fac).start()
                   if re.search(r'(?im)^#+.*' + re.escape(a), fac) else -1) for a in ANCHORS}
        if pos.get("Data Deep Dive", -1) > 0:
            for a in ("Anchor 2", "Anchor 3"):
                if pos.get(a, -1) > pos["Data Deep Dive"]:
                    R.blocker('T1-13', f"MCS-14 — {a} appears AFTER the Data Deep Dive. The whole "
                                       f"pedagogical move is that the question is defined before the numbers.")
        # MCS-15 twist
        if not re.search(r'(?im)^#+.*Twist', fac):
            R.blocker('T1-13', "MCS-15 — no Twist section and no explicit 'Twist: N/A' justification")
        elif re.search(r'(?i)twist\s*[:—-]\s*N/?A', fac):
            R.gap('T1-13', "MCS-15 — Twist is N/A; confirm the justification is written out")
        # MCS-16 timed plan
        plan = re.search(r'(?is)45[\s-]*minute Socratic plan(.*?)(?=^#+\s|\Z)', fac, re.M)
        if not plan:
            R.blocker('T1-13', "MCS-16 — no timed 45-minute Socratic plan (SPEC §7.3)")
        else:
            rows = re.findall(r'\|\s*\**(\d{1,2}:\d{2})', plan.group(1))
            if len(rows) < 6:
                R.gap('T1-13', f"MCS-16 — timed plan has {len(rows)} timed rows; SPEC §7.3 shows 7")
            if not re.search(r'(?i)\bTWIST\b', plan.group(1)):
                R.gap('T1-13', "MCS-16 — timed plan never names the moment the twist drops")
    if re.search(r'(?i)\b(death|dying|cancer|casualt|foreclos|evict|bankrupt|terminal|hospice|layoff|'
                 r'unvaccinated|outbreak)\w*', md) and not (
            c.get('sensitivity_note') or (c.get('facilitator') or {}).get('sensitivity_note')):
        R.blocker('T1-13', "SPEC §6 — case touches illness/death/money stress/family hardship but "
                           "case.json has no one-line facilitator sensitivity note")
    R.ok('T1-13')

    # ---- universe of legitimate numbers ------------------------------------
    universe = numbers_in(md) | numbers_in(json.dumps(c))

    # ---- T1-14 CONTINUITY: FACILITATOR.md vs case.md ------------------------
    # This is the Murder Mystery bug, transposed: the killer changed and three
    # documents kept convicting the old suspect.  Here: a number changes in
    # case.md and FACILITATOR.md quietly keeps teaching the old one.
    if fac:
        # (a) keyed cross-check on derived container rows
        for row in ((c.get('derived') or {}).get('containers') or []):
            name = row.get('name', '')
            # try the full name, then progressively shorter model-code tails —
            # FACILITATOR.md writes "RCW4" where case.json writes "B Medical RCW4"
            toks = name.split('(')[0].split()
            keys = [' '.join(toks[i:]) for i in range(len(toks))]
            frow = None
            for k in keys:
                for line in fac.splitlines():
                    if '|' in line and re.search(re.escape(k), line):
                        frow = line; break
                if frow:
                    break
            if frow is None:
                R.gap('T1-14', f"CONTINUITY — derived container '{name}' has no row in FACILITATOR.md")
                continue
            fnums = numbers_in(frow)
            for field in ('volume_cm3', 'surface_area_cm2', 'sa_v', 'cold_life_h'):
                v = row.get(field)
                if v is None:
                    continue
                if not in_universe(float(v), fnums):
                    R.blocker('T1-14', f"CONTINUITY DRIFT — case.json derived {name}.{field} = {v}, but "
                                       f"the FACILITATOR.md row for {name} does not carry that value. "
                                       f"Row reads: {frow.strip()[:120]}")
        # (b) the balance numbers the guide teaches must be the balance numbers of record
        for k in ('A_in_spec_doses', 'B_in_spec_doses', 'margin_pct'):
            v = bal.get(k)
            if v is None:
                continue
            if not in_universe(float(v), numbers_in(fac)):
                R.blocker('T1-14', f"CONTINUITY DRIFT — balance_test.{k} = {v} appears nowhere in "
                                   f"FACILITATOR.md. The guide is running a different arithmetic than "
                                   f"the case of record.")
        # (c) sweep: guide numbers with no home in case.md or case.json
        fac_body = re.sub(r'(?is)45[\s-]*minute Socratic plan.*?(?=^#+\s|\Z)', '', fac, flags=re.M)
        fac_body = re.sub(r'\d{1,2}:\d{2}', ' ', fac_body)          # clock times
        fac_body = re.sub(r'(?i)~?\d+[\s-]*(?:minute|second|min|sec)s?\b', ' ', fac_body)  # facilitation timings
        orphan_nums = sorted({v for v in numbers_in(fac_body) if not in_universe(v, universe)})
        if orphan_nums:
            R.gap('T1-14', f"CONTINUITY — {len(orphan_nums)} number(s) in FACILITATOR.md appear in "
                           f"neither case.md nor case.json: {orphan_nums[:20]}. Each is either a "
                           f"guide-only derived figure (fine, but it will drift silently) or a stale "
                           f"value. Move derived figures into case.json.derived so they are checkable.")
    R.ok('T1-14')

    # ---- T1-15 CONTINUITY: published artifacts ------------------------------
    pub = os.path.join(cdir, 'published')
    if os.path.isdir(pub):
        src_mtime = os.path.getmtime(md_path)
        fac_mtime = os.path.getmtime(fac_path) if os.path.exists(fac_path) else 0
        TOL = 120                              # seconds — a build takes time
        for fn in sorted(os.listdir(pub)):
            fp = os.path.join(pub, fn)
            if fn.startswith('~$') or not os.path.isfile(fp):
                continue
            # pair each artifact with the source it is generated FROM
            guideish = bool(re.search(r'(?i)guide|facilitator', fn))
            src_name = 'FACILITATOR.md' if guideish else 'case.md'
            src_t = fac_mtime if guideish else src_mtime
            lag = src_t - os.path.getmtime(fp)
            if lag > TOL:
                R.blocker('T1-15', f"CONTINUITY STALE — published/{fn} was built {int(lag/60)} min "
                                   f"BEFORE the last edit to its source ({src_name}). The packet a "
                                   f"student downloads no longer matches the case of record. "
                                   f"Re-run case-publish.")
            if fn.lower().endswith('.xlsx'):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(fp, data_only=True)
                    cells = []
                    for ws in wb.worksheets:
                        for r in ws.iter_rows():
                            for cell in r:
                                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                                    cells.append((ws.title, cell.coordinate, float(cell.value)))
                                elif isinstance(cell.value, str):
                                    for v in numbers_in(cell.value):
                                        cells.append((ws.title, cell.coordinate, v))
                    bad = [(s, co, v) for s, co, v in cells if not in_universe(v, universe)]
                    if bad:
                        R.gap('T1-15', f"CONTINUITY — {len(bad)} numeric cell(s) in published/{fn} have no "
                                       f"match in case.md or case.json, e.g. "
                                       f"{[f'{s}!{co}={v}' for s, co, v in bad[:8]]}")
                except ImportError:
                    R.gap('T1-15', f"openpyxl not installed; could not continuity-check {fn}")
            if fn.lower().endswith('.docx'):
                try:
                    import docx
                    d = docx.Document(fp)
                    txt = "\n".join(p.text for p in d.paragraphs)
                    for t in d.tables:
                        for r in t.rows:
                            txt += "\n" + " | ".join(cell.text for cell in r.cells)
                    dnums = numbers_in(txt)
                    missing = sorted({v for v in dnums if not in_universe(v, universe)})
                    if missing:
                        R.gap('T1-15', f"CONTINUITY — {len(missing)} number(s) in published/{fn} appear in "
                                       f"neither case.md nor case.json: {missing[:12]}")
                    # the other direction: an Appendix A datum that never made the docx
                    for row in ((c.get('derived') or {}).get('containers') or []):
                        v = row.get('cold_life_h')
                        if v is not None and not in_universe(float(v), dnums):
                            R.blocker('T1-15', f"CONTINUITY DRIFT — published/{fn} is missing the cold "
                                               f"life {v} h for {row.get('name')}; the student packet and "
                                               f"case.md disagree on the data.")
                except ImportError:
                    R.gap('T1-15', f"python-docx not installed; could not continuity-check {fn}")
    else:
        R.polish('T1-15', "no published/ directory yet — nothing to drift")
    R.ok('T1-15')

    # ---- T1-16 QA hash currency (MCS-11) ------------------------------------
    h = hashlib.sha256(open(md_path, 'rb').read()).hexdigest()
    q = c.get('qa') or {}
    recorded = q.get('case_md_sha256')
    if not recorded:
        R.gap('T1-16', "MCS-11 — no case_md_sha256 recorded yet (expected on a first run)")
    elif recorded != h:
        R.blocker('T1-16', f"MCS-11 — QA hash is STALE. recorded {recorded[:16]}..., actual {h[:16]}.... "
                           f"case.md changed after the last gate run; every finding below may describe a "
                           f"different file.")
    R.ok('T1-16')

    # ---- T1-17 cohort-size bands (MCS-25) -----------------------------------
    bands = c.get('cohort_bands')
    if not bands:
        R.gap('T1-17', "MCS-25 — case.json declares no cohort_bands[]. A Socratic discussion that works "
                       "with 12 does not work with 24. Declare the bands the case is tuned for, e.g. "
                       "[{\"band\":\"8-12\",\"note\":\"...\"},{\"band\":\"13-18\",...}], and say what "
                       "changes in each.")
    else:
        for b in bands:
            if not b.get('band') or not b.get('note'):
                R.blocker('T1-17', f"MCS-25 — cohort band {b} needs both a size range and a note on what "
                                   f"changes at that size")
    R.ok('T1-17')

    # ---- T1-18 interactive HTML present, fresh, leak-free (MCS-37) ----------
    if os.path.isdir(pub):
        html_files = [fn for fn in os.listdir(pub) if fn.lower().endswith('.html')]
        if not html_files:
            R.blocker('T1-18', "MCS-37 — published/ exists but ships no interactive .html. Every "
                               "published case needs a self-contained interactive artifact "
                               "alongside its docx/xlsx packets, not just narrative documents.")
        else:
            for fn in html_files:
                fp = os.path.join(pub, fn)
                txt = open(fp, encoding='utf-8', errors='ignore').read()
                lag = src_mtime - os.path.getmtime(fp)
                if lag > TOL:
                    R.blocker('T1-18', f"CONTINUITY STALE — published/{fn} was built {int(lag/60)} "
                                       f"min BEFORE the last edit to case.md. Rebuild it.")
                # A calculator's own arithmetic constants (params, circuit data) may legitimately
                # repeat case data. What must NEVER appear, in script or in prose, is the
                # balance_test verdict itself printed as a static literal -- that is MCS-37's
                # actual teeth. Same leak pattern as tier2_probes.py probe_cold_solve.
                # Decorative <style>/<svg> blocks are excluded from the scan: a small round
                # margin_pct (e.g. 5.0) collides constantly with CSS pixel values and SVG path
                # coordinates, which can never be "the answer a fresh reader could just read
                # off the page" -- <script> stays IN scope, since a hardcoded literal there is
                # exactly the leak this check exists to catch.
                leak_scan_txt = re.sub(r'(?is)<style.*?</style>|<svg\b.*?</svg>', ' ', txt)
                bal = c.get('balance') or c.get('balance_test') or {}
                leaks = []
                for k in ('A_in_spec_doses', 'B_in_spec_doses', 'margin_pct'):
                    v = bal.get(k)
                    if v is None:
                        continue
                    # margin_pct is inherently a percent/points figure. A small round value
                    # (5, 10, 15...) collides constantly with unrelated case vocabulary --
                    # "Year 5", "5-year", a <select> option's value="5" -- that share nothing
                    # with the balance verdict except being the same small integer. The real
                    # leak signature this check exists to catch is the value stated AS a
                    # percent/points figure (e.g. "5%", "5.0 percentage points"), so for
                    # margin_pct only, require a %/point unit within a short window after the
                    # match. Other keys keep the original bare-token match.
                    require_unit = (k == 'margin_pct')
                    for form in (f"{v:,}", str(v), str(int(v)) if float(v).is_integer() else str(v)):
                        pat = r'(?<![\w.])' + re.escape(form) + r'(?![\w-])'
                        for m in re.finditer(pat, leak_scan_txt):
                            if require_unit:
                                window = leak_scan_txt[m.end():m.end() + 15]
                                if not re.search(r'\s*%|\s*point', window, re.I):
                                    continue
                            leaks.append(f"balance_test.{k} = {v} appears in {fn}")
                            break
                        if leaks and leaks[-1].startswith(f"balance_test.{k}"):
                            break
                if leaks:
                    R.blocker('T1-18', f"MCS-37 — interactive artifact leaks the answer: {leaks}. "
                                       f"A self-checking calculator must derive its own arithmetic "
                                       f"live and never print the case's balance_test verdict as a "
                                       f"static literal a fresh reader could just read off the page.")
                # Everything outside <script>...</script> (prose, labels, table cells) should
                # still resolve against the case's own universe of numbers, same drift check
                # every other published packet gets under T1-15.
                script_stripped = re.sub(r'(?is)<script.*?</script>', ' ', txt)
                dnums = numbers_in(script_stripped)
                missing = sorted({v for v in dnums if not in_universe(v, universe)})
                if missing:
                    R.gap('T1-18', f"CONTINUITY — {len(missing)} number(s) in {fn}'s prose/markup "
                                   f"(outside <script>) appear in neither case.md nor case.json: "
                                   f"{missing[:12]}")
        R.ok('T1-18')
    else:
        R.polish('T1-18', "no published/ directory yet — MCS-37 has nothing to check")

    # ---- T1-19 hover+hyperlink math glossary, verified domains (MCS-39) -----
    # Two authoring styles are legal: (a) hand-written static anchors
    # <a class="gterm" href="...">, or (b) a JS data object of {label: {def, url}}
    # entries rendered into the page at load time (Cold Chain's pattern -- the
    # anchors themselves only exist after the browser runs renderGlossary(), so
    # a static-text scan of class="gterm" alone would under-count real coverage
    # to just the one template literal that builds them). Count style (b) entries
    # directly from the data so JS-rendered glossaries aren't penalized for not
    # having pre-baked markup.
    ALLOWLIST_DOMAINS = ('khanacademy.org', 'mathsisfun.com', 'cuemath.com',
                         'purplemath.com', 'ck12.org', 'nctm.org', 'openstax.org')
    MIN_GLOSSARY_TERMS = 8
    if os.path.isdir(pub):
        html_files = [fn for fn in os.listdir(pub) if fn.lower().endswith('.html')]
        for fn in html_files:
            txt = open(os.path.join(pub, fn), encoding='utf-8', errors='ignore').read()

            static_terms = re.findall(r'class="gterm"[^>]*href="([^"]+)"', txt)
            static_terms += re.findall(r'href="([^"]+)"[^>]*class="gterm"', txt)
            static_terms = [u for u in static_terms if '${' not in u]  # unresolved template, not a real link

            data_entries = re.findall(
                r'"[A-Za-z][A-Za-z \-]*":\s*\{def:\s*"(?:[^"\\]|\\.)*"\s*,\s*url:\s*"(https://[^"]+)"',
                txt)

            # If a JS data object is present, it's the source of truth for both the
            # count and the URLs to validate -- any co-occurring static class="gterm"
            # match is almost always just the unresolved template scaffolding (or an
            # unrelated anchor), not a second independent set of real links.
            all_urls = data_entries if data_entries else static_terms
            n_terms = len(data_entries) if data_entries else len(static_terms)
            has_wiring = bool(re.search(r'class="gterm"', txt))

            if not has_wiring:
                R.blocker('T1-19', f"MCS-39 — {fn} has no class=\"gterm\" hover/hyperlink markup at "
                                   f"all (static or templated). Every math vocabulary term the case "
                                   f"actually uses should be hoverable (a plain-English definition) "
                                   f"AND clickable (a real external lesson) -- covering the math, not "
                                   f"just the story.")
            elif n_terms < MIN_GLOSSARY_TERMS:
                R.blocker('T1-19', f"MCS-39 — {fn} defines only {n_terms} glossary term(s) with a "
                                   f"real url, needs >= {MIN_GLOSSARY_TERMS}. Presence of the "
                                   f"class=\"gterm\" template alone isn't enough -- the term coverage "
                                   f"itself is thin.")
            else:
                bad = [u for u in all_urls if not (u.startswith('https://')
                                                    and any(d in u for d in ALLOWLIST_DOMAINS))]
                if bad:
                    R.blocker('T1-19', f"MCS-39 — {fn} has {len(bad)} glossary link(s) that are not "
                                       f"https:// to an allowlisted educational domain "
                                       f"{ALLOWLIST_DOMAINS}: {bad[:5]}. No URL from memory (MCS-21) "
                                       f"applies here too -- every link must have actually been "
                                       f"fetched and confirmed before shipping, not guessed.")
        R.ok('T1-19')
    else:
        R.polish('T1-19', "no published/ directory yet — MCS-39 has nothing to check")

    return R, h


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--case', required=True)
    ap.add_argument('--project', default=os.getcwd())
    ap.add_argument('--json', default=None)
    a = ap.parse_args()
    R, h = run(a.project, a.case)

    print(f"TIER 1 — STRUCTURAL   case: {a.case}   sha256(case.md)={h[:16]}")
    print(f"checks run: {len(R.checks)}   "
          f"PASS {sum(1 for v in R.checks.values() if v == 'PASS')} · "
          f"GAP {sum(1 for v in R.checks.values() if v == 'GAP')} · "
          f"FAIL {sum(1 for v in R.checks.values() if v == 'FAIL')}")
    for cid in sorted(R.checks):
        print(f"  {cid}  {R.checks[cid]}")
    for lvl, items in (("BLOCKER", R.B), ("GAP", R.G), ("POLISH", R.P)):
        print(f"\n=== {lvl} ({len(items)}) ===")
        for i in items:
            print(" *", i)
    verdict = 'TIER1_PASS' if not R.B else 'BLOCKED'
    print(f"\nVERDICT: {verdict}  ({len(R.B)} blockers, {len(R.G)} gaps)")
    if a.json:
        json.dump({"case": a.case, "sha256": h, "verdict": verdict, "checks": R.checks,
                   "blockers": R.B, "gaps": R.G, "polish": R.P},
                  open(a.json, 'w'), indent=1)
    sys.exit(0 if not R.B else 1)


if __name__ == '__main__':
    main()
