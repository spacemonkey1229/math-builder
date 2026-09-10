import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = "1A3A5C"
wb = openpyxl.Workbook()

def style_header(ws, row=1, ncols=6):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def set_font(ws, rows, cols, name="Arial", size=10):
    for r in rows:
        for c in cols:
            ws.cell(row=r, column=c).font = Font(name=name, size=size)

# ---------------- README ----------------
ws = wb.active
ws.title = "README"
ws["A1"] = "The Weekly Rate — Data Workbook"
ws["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)
readme_rows = [
    ("", ""),
    ("Case", "The Weekly Rate (cases/the-weekly-rate)"),
    ("Purpose", "Raw, cited data for the Ownership-vs-Subscription e-bike decision Dara Voss faces."),
    ("", ""),
    ("Sheet", "Maps to"),
    ("Appendix A — Cost Components", "case.md Appendix A (S1-S4)"),
    ("Appendix B — Distractor Plan", "case.md Appendix B (S5b)"),
    ("Appendix C — Modeling Assumptions", "case.md Appendix C (MA1-MA4)"),
    ("Work Area", "Blank — for your own D1-D6 calculations. Nothing here is pre-computed."),
    ("", ""),
    ("Note", "This workbook gives you the DATA, not the ANSWER. The cost-function derivation, the "
             "break-even week, the x-intercept, and the percent margin are your own work — see "
             "Challenge 3 (D1-D6) in the case study."),
]
r = 2
for a, b in readme_rows:
    ws.cell(row=r, column=1, value=a).font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=r, column=2, value=b).font = Font(name="Arial", size=10)
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
    r += 1
autosize(ws, [26, 90])

# ---------------- Appendix A ----------------
ws = wb.create_sheet("Appendix A - Cost Components")
ws.append(["Component", "Value", "Units", "Source (case.json sources[].id)"])
rows = [
    ["Zoomo Standard Plan weekly price (all-inclusive)", 39.00, "USD/week", "S1"],
    ["RadRunner 3 Plus purchase price", 2299.00, "USD", "S2"],
    ["RadRunner 3 Plus battery capacity", 672, "Wh", "S2"],
    ["RadRunner 3 Plus manufacturer range, low end", 25, "miles/charge", "S2"],
    ["RadRunner 3 Plus manufacturer range, high end", 45, "miles/charge", "S2"],
    ["U.S. average residential electricity price, June 2026", 0.1834, "USD/kWh", "S3"],
    ["Annual e-bike maintenance, casual rider (upper bound)", 200, "USD/year", "S4"],
    ["Annual e-bike maintenance, regular commuter (lower bound)", 400, "USD/year", "S4"],
    ["Electricity cost, 2,000-5,000 mi/yr rider, low end", 10, "USD/year", "S4"],
    ["Electricity cost, 2,000-5,000 mi/yr rider, high end", 30, "USD/year", "S4"],
]
for row in rows:
    ws.append(row)
style_header(ws, ncols=4)
set_font(ws, range(2, len(rows) + 2), range(1, 5))
autosize(ws, [50, 12, 16, 32])

# ---------------- Appendix B ----------------
ws = wb.create_sheet("Appendix B - Distractor Plan")
ws.append(["Whizz Storm-1 weekly plan (NOT Dara's choice)", "Value", "Units", "Source"])
rows = [
    ["Weekly rate", 69.00, "USD/week", "S5b"],
    ["One-time mandatory deposit", 99.00, "USD", "S5b"],
    ["Total paid at week 0 (deposit only)", 99.00, "USD", "derived"],
    ["Total paid at week 4", 375.00, "USD", "derived (99 + 69*4)"],
]
for row in rows:
    ws.append(row)
style_header(ws, ncols=4)
set_font(ws, range(2, len(rows) + 2), range(1, 5))
autosize(ws, [42, 12, 14, 26])

# ---------------- Appendix C ----------------
ws = wb.create_sheet("Appendix C - Assumptions")
ws.append(["#", "Assumption", "Value used", "Basis"])
rows = [
    ["MA1", "Dara's average weekly mileage", "200 mi/week", "6 shifts/week at ~33 mi/shift"],
    ["MA2", "Charging efficiency used", "25 mi/charge (low end)", "Loaded delivery bike, frequent stops"],
    ["MA3", "Dara's own annual maintenance if she owns", "$900.00/year", "Scaled up from S4's commuter figure"],
    ["MA4", "Dara's decision horizon", "91 weeks (21 months)", "Length of her certificate program"],
]
for row in rows:
    ws.append(row)
style_header(ws, ncols=4)
set_font(ws, range(2, len(rows) + 2), range(1, 5))
autosize(ws, [6, 40, 24, 46])

# ---------------- Work Area ----------------
ws = wb.create_sheet("Work Area")
ws["A1"] = "Work Area — your own calculations (nothing pre-filled)"
ws["A1"].font = Font(name="Arial", bold=True, size=12, color=NAVY)
labels = [
    "D1: Option A weekly slope (USD/week) =",
    "D1: Option A y-intercept (USD) =",
    "D1: Option B weekly slope (USD/week) =",
    "D1: Option B y-intercept (USD) =",
    "D2: Break-even week (solve C_A(w) = C_B(w)) =",
    "D3: x-intercept of S(w) = C_B(w) - C_A(w) =",
    "D3: y-intercept of S(w), i.e. S(0) =",
    "D4: Whizz plan slope (from Appendix B graph) =",
    "D4: Whizz plan y-intercept (from Appendix B graph) =",
    "D6: Option A total cost at w = 91 =",
    "D6: Option B total cost at w = 91 =",
    "D6: Dollar difference =",
    "D6: Percent margin of the more expensive option =",
]
r = 3
for lab in labels:
    ws.cell(row=r, column=1, value=lab).font = Font(name="Arial", size=10)
    ws.cell(row=r, column=2).font = Font(name="Arial", size=10)
    ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="FFF8E1")
    r += 1
autosize(ws, [46, 18])

wb.save("The Weekly Rate - Data.xlsx")
print("wrote xlsx")
