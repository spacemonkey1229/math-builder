#!/usr/bin/env python3
"""Build the Data workbook for The Dosing Curve. Raw appendix data only --
never pre-computes D1-D4 (Riemann sums, antiderivatives/FTC, fAUC/MIC ratios,
or per-animal cost); students do that work in the Work Area sheet."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARIAL = "Arial"

def style_title(ws, cell, size=14):
    ws[cell].font = Font(name=ARIAL, size=size, bold=True)

def style_sub(ws, cell):
    ws[cell].font = Font(name=ARIAL, size=11, italic=True)

def style_header_row(ws, row, ncols):
    fill = PatternFill(fill_type="solid", fgColor="DDDDDD")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=ARIAL, bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

def set_body_font(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and (cell.font is None or cell.font.name != ARIAL):
                cell.font = Font(name=ARIAL, size=11)

def col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = openpyxl.Workbook()

# ---------- README ----------
ws = wb.active
ws.title = "README"
ws["A1"] = "The Dosing Curve — Data Workbook"
style_title(ws, "A1")
ws["A2"] = "MS | Math | Case Study | 2026 | Adventure"
style_sub(ws, "A2")
ws["A5"] = "What this is"
ws["B5"] = ("Raw appendix data from The Dosing Curve case study, laid out for you to compute "
            "over. Nothing here has been pre-solved -- the sheets give you the sourced "
            "pharmacokinetic and cost figures the case cites; the Riemann sums, antiderivatives, "
            "definite integrals, fAUC/MIC ratios, and per-animal cost comparison (D1-D4 in the "
            "case study) are your work to do.")
ws["A6"] = "Case"
ws["B6"] = ("cases/the-dosing-curve/case.md — see the case study docx for the full Overview, "
            "Background, Dilemma, and Challenges.")
ws["A8"] = "Sheet"
ws["B8"] = "Maps to"
ws["A9"] = "Appendix A — PK Parameters"
ws["B9"] = "Appendix A: single-dose pharmacokinetic summary statistics, conventional vs. long-acting. Source: [S1]."
ws["A10"] = "Appendix B — Concentration-Time Data"
ws["B10"] = "Appendix B: sampled/modeled concentration-time table for both formulations, the raw points D1's Riemann sum needs. Source: [S1], reconstructed per MA1."
ws["A11"] = "Appendix C — Efficacy Target"
ws["B11"] = "Appendix C: fAUC24h/MIC efficacy target and PK/PD cutoff concentrations, plus the CLSI breakpoint. Sources: [S2], [S3]."
ws["A12"] = "Appendix D — Dosing Regimens"
ws["B12"] = "Appendix D: both product labels' dose rate, frequency, and pre-slaughter withdrawal. Sources: [S4], [S5], [S6]."
ws["A13"] = "Appendix E — Cost Economics"
ws["B13"] = "Appendix E: per-mL drug price, injection-site-lesion cost and incidence rates, and the operation's budget cap. Sources: [S7], [S8], [S9]; MA5."
ws["A14"] = "Appendix F — Modeling Assumptions"
ws["B14"] = "Appendix F: MA1-MA5, every stated (not sourced) figure used to bridge real data to this scenario."
ws["A15"] = "D1-D4 Work Area"
ws["B15"] = ("A blank area for your own Riemann-sum table, fitted-curve antiderivative/FTC "
             "evaluation, fAUC/MIC comparison, and per-animal cost comparison. This is the "
             "assignment -- it is intentionally empty.")
for r in (5, 6, 8, 9, 10, 11, 12, 13, 14, 15):
    ws[f"A{r}"].font = Font(name=ARIAL, bold=(r == 8))
    ws[f"B{r}"].font = Font(name=ARIAL)
    ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
col_widths(ws, [28, 95])
ws.row_dimensions[5].height = 55
ws.row_dimensions[6].height = 30
for r in (9, 10, 11, 12, 13):
    ws.row_dimensions[r].height = 30
ws.row_dimensions[15].height = 55

# ---------- Appendix A ----------
ws = wb.create_sheet("Appendix A — PK Parameters")
rows = [
    ("Parameter", "Conventional IM", "Long-Acting IM"),
    ("Peak concentration, Cmax (µg/mL)", "13.57 ± 5.83", "8.72 ± 2.47"),
    ("Time to peak, Tmax (h)", "0.46 ± 0.09", "0.60 ± 0.28"),
    ("Half-life (h)", "10.47 ± 1.30", "27.96 ± 11.66"),
    ("Single-dose AUC (µg·h/mL)", "80.21 ± 12.71", "87.35 ± 10.91"),
    ("Bioavailability (%)", 69.71, 83.15),
]
for r in rows:
    ws.append(r)
style_header_row(ws, 1, 3)
set_body_font(ws)
col_widths(ws, [40, 24, 24])
ws.append([])
ws.append(["Source: [S1] Aktaş İ, Yarsan E. Pharmacokinetics of Conventional and Long-Acting Oxytetracycline Preparations in Kilis Goat. Frontiers in Veterinary Science 4:229 (2017). Single-dose 20 mg/kg IM in goats -- species and dose differ from this case's cattle scenario; the best available real comparative dataset for the two formulation types.", None, None])
ws.cell(row=ws.max_row, column=1).font = Font(name=ARIAL, italic=True, size=10)
ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[ws.max_row].height = 45

# ---------- Appendix B ----------
ws = wb.create_sheet("Appendix B — Concentration-Time")
ws.append(["Time (h)", "Conventional IM, modeled C (µg/mL)", "Long-Acting IM, modeled C (µg/mL)"])
data_b = [
    (0.25, 4.93, 1.98), (0.5, 5.15, 2.13), (1, 5.00, 2.12), (2, 4.68, 2.07),
    (4, 4.10, 1.97), (8, 3.15, 1.78), (12, 2.42, 1.61), (24, 1.09, 1.20),
    (36, 0.49, 0.89), (48, 0.22, 0.66), (60, 0.10, 0.49), (72, 0.05, 0.36), (96, 0.01, 0.20),
]
for r in data_b:
    ws.append(r)
style_header_row(ws, 1, 3)
set_body_font(ws)
col_widths(ws, [12, 34, 34])
ws.append([])
ws.append(["Disclosed limitation: [S1] sampled blood at these 13 real timepoints but published only summary statistics (Appendix A) and a plotted figure -- no raw concentration table. This table is a reconstructed curve (MA1), not measured data. Time = 0 (moment of injection) is not shown here but is C = 0 by definition, needed for a Riemann sum starting at t = 0.", None, None])
ws.cell(row=ws.max_row, column=1).font = Font(name=ARIAL, italic=True, size=10)
ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[ws.max_row].height = 45

# ---------- Appendix C ----------
ws = wb.create_sheet("Appendix C — Efficacy Target")
ws.append(["Quantity", "Value"])
c_rows = [
    ("Target fAUC24h/MIC ratio (efficacy cutoff)", "19.1–28 h; fixed at 24 h for a 24-h dosing interval"),
    ("PK/PD cutoff concentration, calves (mg/L)", 1.0),
    ("PK/PD cutoff concentration, adult cattle (mg/L)", 2.0),
    ("CLSI VET01S susceptible breakpoint, oxytetracycline vs. M. haemolytica (mg/L)", "≤2"),
    ("CLSI VET01S resistant breakpoint, oxytetracycline vs. M. haemolytica (mg/L)", "≥8"),
]
for r in c_rows:
    ws.append(r)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [58, 40])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.append([])
ws.append(["Sources: [S2] Winter EA, et al. Frontiers in Microbiology (2024) -- PK/PD cutoffs derived for a long-acting-type regimen (20 mg/kg IM, 48-h interval); [S3] CLSI VET01S, via EMA Q&A secondary citation.", None])
ws.cell(row=ws.max_row, column=1).font = Font(name=ARIAL, italic=True, size=10)
ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[ws.max_row].height = 45

# ---------- Appendix D ----------
ws = wb.create_sheet("Appendix D — Dosing Regimens")
ws.append(["Quantity", "The Single Shot (Long-Acting)", "The Four-Day Course (Conventional)"])
d_rows = [
    ("Labeled dose (mg oxytetracycline / lb body weight)", "9 mg/lb, single dose", "3–5 mg/lb/day"),
    ("Dosing frequency", "Single IM or SC injection", "Once daily, not to exceed 4 consecutive days"),
    ("Pre-slaughter withdrawal period (days)", 28, 28),
]
for r in d_rows:
    ws.append(r)
style_header_row(ws, 1, 3)
set_body_font(ws)
col_widths(ws, [46, 34, 34])
for r in range(2, ws.max_row + 1):
    for c in (1, 2, 3):
        ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
ws.append([])
ws.append(["Sources: [S4] DailyMed -- Liquamycin LA-200 FDA label; [S5] DailyMed -- Oxytetracycline 200 FDA label. Historical note [S6]: Xia, Nielsen & Gyrd-Hansen 1983 -- plasma >0.5 µg/mL maintained ~60h by either prep after 20 mg/kg IM, but pronounced injection-site tissue damage 1-2 weeks post-injection for long-acting vs. minimal for conventional.", None, None])
ws.cell(row=ws.max_row, column=1).font = Font(name=ARIAL, italic=True, size=10)
ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[ws.max_row].height = 55

# ---------- Appendix E ----------
ws = wb.create_sheet("Appendix E — Cost Economics")
ws.append(["Quantity", "Value"])
e_rows = [
    ("Per-mL drug price, long-acting formulation (USD)", 0.31998),
    ("Per-mL drug price, conventional formulation (USD)", 0.11598),
    ("Both formulations, concentration (mg/mL)", 200),
    ("Industry-documented injection-site-lesion cost (USD per head marketed, 1995 audit)", 7.05),
    ("Injection-site-lesion incidence rate, 1995 survey", "11%"),
    ("Injection-site-lesion incidence rate, most recent survey (fed cattle, 2022)", "0.2%"),
    ("Operation's fixed per-head antimicrobial treatment budget (USD)", 25),
]
for r in e_rows:
    ws.append(r)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [64, 24])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.append([])
ws.append(["These per-mL prices are derived (not directly published) from [S7]'s same-retailer list pricing for each product's 500 mL bottle, at each formulation's stated 200 mg/mL concentration. Converting a per-mL price into a per-dose or per-course cost still requires the animal's body weight (MA3) and, for the conventional course, a specific point within the labeled 3-5 mg/lb/day range (MA4) -- neither is computed for you here.", None])
ws.cell(row=ws.max_row, column=1).font = Font(name=ARIAL, italic=True, size=10)
ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[ws.max_row].height = 55
ws.append(["Sources: [S7] EntirelyPets Rx retail list pricing; [S8] 1995 National Beef Quality Audit via NDSU Extension; [S9] National Beef Quality Audit-2022 via SDSU Extension. Budget cap is Modeling Assumption MA5 (see Appendix F), informed by [S10].", None])
ws.cell(row=ws.max_row, column=1).font = Font(name=ARIAL, italic=True, size=10)
ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[ws.max_row].height = 30

# ---------- Appendix F ----------
ws = wb.create_sheet("Appendix F — Assumptions")
ws.append(["#", "Assumption", "Value used", "Basis"])
f_rows = [
    ("MA1", "Reconstructed concentration-time curve (Appendix B)",
     "One-compartment, first-order-absorption model; elimination rate fixed from Appendix A's real half-life; absorption rate matched to Appendix A's real Tmax; total area fixed to Appendix A's real AUC.",
     "[S1] published only summary PK statistics plus the real 13-point sampling schedule, not a raw concentration table."),
    ("MA2", "Multi-dose superposition, Four-Day Course",
     "Linear, time-invariant PK; cumulative curve = 4 identical single-dose curves, each shifted 24h.",
     "Standard PK principle for a drug with linear elimination at labeled doses."),
    ("MA3", "Representative treated-animal body weight", "500 lb",
     "No specific weight is stated in the case; a commonly cited representative weight for incoming stocker/feeder calves."),
    ("MA4", "Conventional dosing rate used for cost calculation", "5 mg/lb/day (labeled range's upper end)",
     "The label gives a range (3-5 mg/lb/day), not a single number; this case uses the range's upper end."),
    ("MA5", "Operation's fixed per-head antimicrobial treatment budget", "$25 per head",
     "Set close to a real USDA APHIS NAHMS Feedlot-survey average BRD treatment cost benchmark ([S10])."),
]
for r in f_rows:
    ws.append(r)
style_header_row(ws, 1, 4)
set_body_font(ws)
col_widths(ws, [8, 40, 48, 56])
for r in range(2, ws.max_row + 1):
    for c in (2, 3, 4):
        ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 55
ws.append([])
ws.append(["These figures are not sourced -- they are stated scenario assumptions chosen by the case author, each with a real basis named above. You may challenge any of them in discussion -- that is a legitimate move.", None, None, None])
ws.cell(row=ws.max_row, column=1).font = Font(name=ARIAL, italic=True, size=10)

# ---------- D1-D4 Work Area ----------
ws = wb.create_sheet("D1-D4 Work Area")
ws["A1"] = "Work Area — D1 through D4"
style_title(ws, "A1", size=13)
ws["A2"] = ("This sheet is intentionally blank except for column headers. Pull your raw inputs "
            "from the Appendix A-F sheets -- do not retype values by hand where you can "
            "reference the cell.")
ws["A2"].font = Font(name=ARIAL, italic=True)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:F2")
ws.row_dimensions[2].height = 40

sections = [
    (4, "D1 — Riemann-sum table for AUC (each option)",
     ["Option", "Subinterval (t_lo–t_hi, h)", "Δt (h)", "C(t_lo), C(t_hi) (µg/mL)", "Trapezoid contribution (µg·h/mL)", "Running total"]),
    (12, "D2 — Fitted C(t), antiderivative (power rule), FTC definite integral",
     ["Option", "Fitted C(t) (state your model)", "General antiderivative F(t) + C", "F(24) − F(0)", "Compared to D1 total", "Work shown"]),
    (19, "D3 — fAUC/MIC comparison",
     ["Option", "Your computed AUC0–24h", "MIC used and why (population)", "fAUC24h/MIC ratio", "Compared to Appendix C's target"]),
    (26, "D4 — Shaded AUC graph and per-animal cost comparison",
     ["Option", "Drug cost (mL × price/mL)", "Handling events", "Expected lesion cost", "Total per-animal cost", "Notes"]),
    (33, "Recommendation",
     ["Which option would you sign?", "Defended using D1-D4", "What would have to change for the recommendation to flip"]),
]
for start_row, title, headers in sections:
    ws.cell(row=start_row, column=1, value=title).font = Font(name=ARIAL, bold=True, size=12)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=start_row + 1, column=c, value=h)
    style_header_row(ws, start_row + 1, len(headers))
    for blank_r in range(start_row + 2, start_row + 5):
        for c in range(1, len(headers) + 1):
            ws.cell(row=blank_r, column=c, value=None)

col_widths(ws, [28, 28, 28, 28, 28, 28])

wb.save("cases/the-dosing-curve/published/The Dosing Curve - Data.xlsx")
print("wrote The Dosing Curve - Data.xlsx")
