#!/usr/bin/env python3
"""Build the Data workbook for The Repeat Match. Raw appendix data only --
never pre-computes D1-D6 or the marker-layout cost comparison; students do
that work in the work-area sheet."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARIAL = "Arial"


def style_title(ws, cell, size=14):
    ws[cell].font = Font(name=ARIAL, size=size, bold=True)


def style_sub(ws, cell):
    ws[cell].font = Font(name=ARIAL, size=11, italic=True)


def style_header_row(ws, row, ncols):
    fill = PatternFill(fill_type="solid", fgColor="DDDDDD")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=ARIAL, bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def set_body_font(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and (cell.font is None or cell.font.name != ARIAL):
                cell.font = Font(name=ARIAL, size=11)


def col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def note(ws, text):
    ws.append([text])
    ws.cell(row=ws.max_row, column=1).font = Font(name=ARIAL, italic=True, size=10)
    ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")


wb = openpyxl.Workbook()

# ---------- README ----------
ws = wb.active
ws.title = "README"
ws["A1"] = "The Repeat Match — Data Workbook"
style_title(ws, "A1")
ws["A2"] = "MS | Math | Case Study | 2026 | Adventure"
style_sub(ws, "A2")
ws["A5"] = "What this is"
ws["B5"] = ("Raw appendix data from The Repeat Match case study, laid out for you to compute "
            "over. Nothing here has been pre-solved -- the sheets give you the sourced figures "
            "and stated modeling assumptions the case cites; the geometry (D1-D5) and the "
            "marker-layout cost comparison (D6) are your work to do.")
ws["A6"] = "Case"
ws["B6"] = ("cases/the-repeat-match/case.md — see the case study docx for the full Overview, "
            "Background, Dilemma, and Challenges.")
ws["A8"] = "Sheet"
ws["B8"] = "Maps to"
mapping = [
    ("Appendix A — Panel Geometry", "Registration-point coordinates, seam axes, and rotation angles for the front and back-yoke panels. Source: MA1-MA6 (stated modeling assumptions -- internal to Yuki's own CAD file, unsourceable outside it)."),
    ("Appendix B — Repeat Range", "The EFI Reggiani UNICA's real sourced adjustable repeat-length range, the master motif's native repeat length, and a second printer's repeat range. Sources: [S1], MA7, MA8."),
    ("Appendix C — Patent Buffer", "US Patent 5,333,111's stated maximum pattern-alignment buffer, as a fraction of one design repeat and in inches. Source: [S2], derived from [S1]."),
    ("Appendix D — Nap Yardage", "The extra yardage a with-nap (same-direction-only) marker layout requires over a without-nap layout. Sources: [S3], [S4], [S5]."),
    ("Appendix E — Marker Efficiency", "Real measured marker material utilization and waste-range data from a 2024 university capstone report, plus one company's disclosed cutting efficiency and a broader literature baseline. Sources: [S6], [S7], [S8], [S9]."),
    ("Appendix F — Run & Cost", "Fabric cost per yard, total production run size, and base yardage per unit. Sources: [S10], MA9, MA10."),
    ("Modeling Assumptions", "Every MA# figure used above, its value, its basis, and its stated sensitivity -- reproduced from case.md so you don't have to cross-reference the docx while you work."),
    ("D1-D6 Work Area", "A blank area for your own coordinate-plane transformations, symmetry analysis, repeat-tiling diagram inputs, congruent-vs-similar/dilation comparison, and the marker-layout cost comparison. This is the assignment -- it is intentionally empty."),
]
r = 9
for label, desc in mapping:
    ws.cell(row=r, column=1, value=label).font = Font(name=ARIAL, bold=True)
    c = ws.cell(row=r, column=2, value=desc)
    c.font = Font(name=ARIAL)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30 if len(desc) < 140 else 45
    r += 1
col_widths(ws, [32, 95])
ws.row_dimensions[5].height = 45
ws.row_dimensions[6].height = 30

# ---------- Appendix A ----------
ws = wb.create_sheet("Appendix A — Panel Geometry")
rows = [
    ("Quantity", "Value"),
    ("Front-panel registration point coordinates (local CAD panel coordinate system)", "(4, 6), in inches, origin at the panel's bottom-left corner [MA1]"),
    ("Front seam axis (center-front zipper seam) orientation", "Vertical -- coincides with the panel's local y-axis [MA2]"),
    ("Marker cutting-angle rotation applied to the front panel", "90° counterclockwise [MA3]"),
    ("Back-yoke registration point coordinates", "(5, 3), in inches, origin at the panel's bottom-left corner [MA4]"),
    ("Back-yoke seam axis orientation", "Horizontal -- coincides with the panel's local x-axis [MA5]"),
    ("Rotation applied to the third, non-mirrored marker panel referenced in D2", "180° [MA6]"),
]
for row in rows:
    ws.append(row)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [58, 55])
for rr in range(2, ws.max_row + 1):
    ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")
ws.append([])
note(ws, "Every value in this appendix is a stated Modeling Assumption -- these figures live only inside Yuki's own CAD file, and no outside source could ever report them. See the Modeling Assumptions sheet for basis/sensitivity.")

# ---------- Appendix B ----------
ws = wb.create_sheet("Appendix B — Repeat Range")
rows = [
    ("Quantity", "Value"),
    ("Primary printer (EFI Reggiani UNICA) minimum adjustable print repeat length", "25 in / 640 mm [S1]"),
    ("Primary printer (EFI Reggiani UNICA) maximum adjustable print repeat length", "40 in / 1018 mm [S1]"),
    ("Master motif's native repeat length, as designed", "30 in / 762 mm [MA7]"),
    ("Second printer's electronically adjustable repeat length range", "20-30 in / 508-762 mm [MA8]"),
]
for row in rows:
    ws.append(row)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [58, 40])
ws.append([])
note(ws, "Source [S1]: EFI -- Reggiani UNICA product page (efi.com/products/textile/efi-reggiani-analog-printers/efi-reggiani-unica-rotary-screen-printer/) and brochure, 2023, HIGH confidence. MA7/MA8 are stated Modeling Assumptions -- see the Modeling Assumptions sheet.")

# ---------- Appendix C ----------
ws = wb.create_sheet("Appendix C — Patent Buffer")
rows = [
    ("Quantity", "Value"),
    ("Patent's stated maximum pattern-alignment buffer, as a fraction of one design repeat", "One-half (1/2) of one fabric design repeat [S2]"),
    ("Buffer size in inches, using Appendix B's repeat-length range", "12.5-20 in (derived: 1/2 x 25 in and 1/2 x 40 in, from [S1] and [S2])"),
]
for row in rows:
    ws.append(row)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [58, 45])
ws.append([])
note(ws, "Source [S2]: US Patent 5,333,111, 'Garment cutting system having computer assisted pattern alignment' -- USPTO / Google Patents (patents.google.com/patent/US5333111A/en), 1994 grant, HIGH confidence.")

# ---------- Appendix D ----------
ws = wb.create_sheet("Appendix D — Nap Yardage")
rows = [
    ("Quantity", "Value"),
    ("Extra yardage required for a with-nap (same-direction-only) layout vs. a without-nap (flip-allowed) layout", "15-25% [S3][S4], consistent with Guideline 4.120's own qualitative confirmation that a with-nap layout 'generally requires' extra fabric [S5]"),
]
for row in rows:
    ws.append(row)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [58, 60])
ws.cell(row=2, column=2).alignment = Alignment(wrap_text=True, vertical="top")
ws.append([])
note(ws, "Sources: [S3] Sewing With Ease, Fabric Yardage Calculator, 2026, MEDIUM. [S4] Canvas Etc., Fabric Yardage Calculator, 2026, MEDIUM. [S5] Sewing & Craft Alliance, Guideline 4.120 'Napped Fabrics' (sewing.org/files/guidelines/4_120_napped_fabrics.pdf), 2006, HIGH for the qualitative claim (no percentage stated in the primary document).")

# ---------- Appendix E ----------
ws = wb.create_sheet("Appendix E — Marker Efficiency")
rows = [
    ("Quantity", "Value"),
    ("Average marker material utilization, measured Vietnam-factory data (Patagonia 2022 seasons)", "82% [S6]"),
    ("Measured marker waste range, same report", "16-21% [S7]"),
    ("686 Technical Apparel's own publicly disclosed cutting efficiency", "87% [S8]"),
    ("Broader literature baseline, pre-consumer cutting waste", "10-15% [S9]"),
]
for row in rows:
    ws.append(row)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [65, 30])
ws.append([])
note(ws, "Sources: [S6]/[S7] UC Santa Barbara Bren School of Environmental Science & Management, 'Examining Cut-and-Sew Textile Waste within the Apparel Supply Chain' (Final Report, March 22, 2024), Figure 6 p.35 / Table 10 p.34, HIGH. [S8] same report p.9, citing '686 Technical Apparel, n.d.', MEDIUM (secondary citation). [S9] same report p.9, citing Rissanen (2005), MEDIUM (secondary citation).")

# ---------- Appendix F ----------
ws = wb.create_sheet("Appendix F — Run & Cost")
rows = [
    ("Quantity", "Value"),
    ("Brushed/napped softshell face fabric cost", "$39.00 / yard [S10]"),
    ("Jacket line total production run size", "5,000 units [MA9]"),
    ("Base fabric yardage required per finished unit (flip-allowed baseline, before any nap penalty)", "3.0 yd / unit [MA10]"),
]
for row in rows:
    ws.append(row)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [70, 30])
ws.append([])
note(ws, "Source [S10]: Discovery Fabrics, Softshells collection (discoveryfabrics.com/collections/softshells), 2026, MEDIUM -- a retail, small-yardage consumer price used as a proxy for a converter's bulk per-yard cost. MA9/MA10 are stated Modeling Assumptions -- see the Modeling Assumptions sheet.")

# ---------- Modeling Assumptions ----------
ws = wb.create_sheet("Modeling Assumptions")
ws["A1"] = "Modeling Assumptions"
style_title(ws, "A1", size=13)
ws["A2"] = "These figures are NOT sourced. They are stated scenario assumptions chosen by the case author. You may challenge them in discussion -- that is a legitimate move."
ws["A2"].font = Font(name=ARIAL, italic=True)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:D2")
ws.row_dimensions[2].height = 30

header_row = 4
ws.cell(row=header_row, column=1, value="#")
ws.cell(row=header_row, column=2, value="Assumption")
ws.cell(row=header_row, column=3, value="Value used")
ws.cell(row=header_row, column=4, value="Sensitivity")
ma_rows = [
    ("MA1", "Front-panel registration point coordinates", "(4, 6) in", "Does not affect the cost dilemma or margin. Affects only the geometry diagrams (D1-D3)."),
    ("MA2", "Front seam axis orientation", "Vertical (y-axis)", "No effect on cost or margin."),
    ("MA3", "Marker cutting-angle rotation, front panel", "90° CCW", "No effect on cost or margin; 90° chosen for MS-level coordinate arithmetic."),
    ("MA4", "Back-yoke registration point coordinates", "(5, 3) in", "No effect on cost or margin."),
    ("MA5", "Back-yoke seam axis orientation", "Horizontal (x-axis)", "No effect on cost or margin."),
    ("MA6", "Rotation on the third, non-mirrored panel (D2)", "180°", "No effect on cost or margin."),
    ("MA7", "Master motif's native repeat length", "30 in / 762 mm", "No effect on cost or margin; sets up the D4/D5 translation and dilation diagrams."),
    ("MA8", "Second printer's adjustable repeat range", "20-30 in / 508-762 mm", "Flips D5's dilation scale factor directly; never changes which marker-layout policy is cheaper."),
    ("MA9", "Jacket line total production run size", "5,000 units", "Cancels out of the cost margin -- multiplies both options' total cost identically."),
    ("MA10", "Base fabric yardage required per finished unit (flip-allowed baseline)", "3.0 yd / unit", "Cancels out of the cost margin for the same reason as MA9."),
]
for row in ma_rows:
    ws.append(row)
style_header_row(ws, header_row, 4)
set_body_font(ws)
col_widths(ws, [7, 50, 25, 60])
for rr in range(header_row + 1, ws.max_row + 1):
    ws.cell(row=rr, column=4).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- D1-D6 Work Area ----------
ws = wb.create_sheet("D1-D6 Work Area")
ws["A1"] = "Work Area — D1 through D6"
style_title(ws, "A1", size=13)
ws["A2"] = ("This sheet is intentionally blank except for column headers. Pull your raw inputs "
            "from the Appendix A-F sheets and the Modeling Assumptions sheet -- do not retype "
            "values by hand where you can reference the cell.")
ws["A2"].font = Font(name=ARIAL, italic=True)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:E2")
ws.row_dimensions[2].height = 40

sections = [
    (4, "D1 — Front-panel registration point through both composition orders",
     ["Step", "Rotate-then-reflect coordinates", "Reflect-then-rotate coordinates", "Notes"]),
    (10, "D2 — Back-yoke panel build + 180° rotation check",
     ["Step", "Back-yoke coordinates", "Third-panel (180°-only) coordinates", "Matching orientation? (Y/N)", "Notes"]),
    (16, "D3 — Symmetry analysis of the master chevron motif",
     ["Candidate axis tested", "Line of symmetry? (Y/N)", "Why / why not"]),
    (21, "D4 — Repeat-tiling diagram inputs",
     ["Printer setting", "Repeat length used (in)", "Translation vector", "Number of repeats across printable width"]),
    (27, "D5 — Congruent-vs-similar / dilation comparison",
     ["Copy", "Congruent to master? (Y/N)", "Similar to master? (Y/N)", "Scale factor (if dilated)", "Work shown"]),
    (33, "D6 — Marker-layout cost comparison and recommendation",
     ["Layout", "Base yardage per unit", "Yardage premium applied", "Total yards for run", "Total cost for run"]),
]
for start_row, title, headers in sections:
    ws.cell(row=start_row, column=1, value=title).font = Font(name=ARIAL, bold=True, size=12)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=start_row + 1, column=c, value=h)
    style_header_row(ws, start_row + 1, len(headers))
    for blank_r in range(start_row + 2, start_row + 4):
        for c in range(1, len(headers) + 1):
            ws.cell(row=blank_r, column=c, value=None)

col_widths(ws, [30, 28, 28, 28, 28])

wb.save("cases/the-repeat-match/published/The Repeat Match - Data.xlsx")
print("wrote The Repeat Match - Data.xlsx")
