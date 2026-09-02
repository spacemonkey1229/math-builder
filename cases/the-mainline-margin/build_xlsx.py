#!/usr/bin/env python3
"""Build the Data workbook for The Mainline Margin. Raw appendix data only --
never pre-computes D1-D6; students do that work in the Work Area sheet."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARIAL = "Arial"

def style_title(ws, cell, size=14):
    ws[cell].font = Font(name=ARIAL, size=size, bold=True)

def style_sub(ws, cell):
    ws[cell].font = Font(name=ARIAL, size=11, italic=True)

def style_header_row(ws, row, ncols):
    fill = PatternFill(fill_type="solid", fgColor="DDDDDD")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=ARIAL, bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

def set_body_font(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and (cell.font is None or cell.font.name != ARIAL):
                cell.font = Font(name=ARIAL, size=11)

def col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = openpyxl.Workbook()

# ---------- README ----------
ws = wb.active
ws.title = "README"
ws["A1"] = "The Mainline Margin — Data Workbook"
style_title(ws, "A1")
ws["A2"] = "MS | Math | Case Study | 2026-27 | Heroes Academy"
style_sub(ws, "A2")
ws["A5"] = "What this is"
ws["B5"] = ("Raw appendix data from The Mainline Margin case study, laid out for you to compute "
            "over. Nothing here has been pre-solved -- D1 through D6 (building each option's total "
            "annualized-cost expression, substituting real values, simplifying, the certification "
            "check, the Hazen-Williams diameter rearrangement, and the break-even calculation) are "
            "your own work to do.")
ws["A6"] = "Case"
ws["B6"] = "cases/the-mainline-margin/case.md — see the case study docx for the full Overview, Background, Dilemma, and Challenges."
ws["A8"] = "Sheet"
ws["B8"] = "Maps to"
ws["A9"] = "Appendix A — Bid Constraints & Field Layout"
ws["B9"] = "Appendix A: design flow rate, run length, acreage, labor rate, competing bid, annual hours. MA1-MA6."
ws["A10"] = "Appendix B — Option A (Lean Main) Spec"
ws["B10"] = "Appendix B: diameter, weight/ft, material cost/ft, Hazen-Williams C. Sources: S3, S6, S7."
ws["A11"] = "Appendix C — Option B (Margin Main) Spec"
ws["B11"] = "Appendix C: diameter, weight/ft, material cost/ft, Hazen-Williams C. Sources: S3, S6, S8."
ws["A12"] = "Appendix D — NRCS Code 430"
ws["B12"] = "Appendix D: velocity ceiling, friction-loss method and formula. Sources: S1, S2."
ws["A13"] = "Appendix E — EQIP Cost-Share Terms"
ws["B13"] = "Appendix E: reimbursement basis, cost-share rate, per-pound rate, forfeiture condition. Sources: S4, S5."
ws["A14"] = "Appendix F — Pump Energy Cost Figures"
ws["B14"] = "Appendix F: fuel type, energy-cost conversion factor, electricity price, BHP formula. Sources: S9, S10, S12, MA7."
ws["A15"] = "Appendix G — Sources"
ws["B15"] = "Appendix G: full citation table for every figure used, S1-S12."
ws["A16"] = "D1-D6 Work Area"
ws["B16"] = ("A blank area for your own cost expressions, substitution table, simplified "
             "expressions, certification check, rearranged Hazen-Williams formula, and break-even "
             "calculation. This is the assignment -- it is intentionally empty.")
for r in (5, 6, 8, 9, 10, 11, 12, 13, 14, 15, 16):
    ws[f"A{r}"].font = Font(name=ARIAL, bold=(r == 8))
    ws[f"B{r}"].font = Font(name=ARIAL)
    ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
col_widths(ws, [30, 100])
ws.row_dimensions[5].height = 55
ws.row_dimensions[6].height = 30
for r in (9, 10, 11, 12, 13, 14, 15):
    ws.row_dimensions[r].height = 30
ws.row_dimensions[16].height = 55

# ---------- Appendix A ----------
ws = wb.create_sheet("Appendix A — Bid Constraints")
ws.append(["Quantity", "Value"])
a_rows = [
    ("Mainline design flow rate required at the field inlet (GPM)", "440 [MA1]"),
    ("Total mainline run length (ft)", "1,320 [MA2]"),
    ("Acreage irrigated by this mainline (acres)", "88 [MA3]"),
    ("Installation labor rate ($ per foot of pipe installed)", "$25.00 [MA4]"),
    ("Competing contractor's total quoted price, out-of-pocket, net of EQIP (USD)", "$24,000 [MA6]"),
    ("Annual hours the irrigation system operates (hr/yr)", "1,200 [MA5]"),
]
for r in a_rows:
    ws.append(r)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [55, 30])

# ---------- Appendix B ----------
ws = wb.create_sheet("Appendix B — Lean Main")
ws.append(["Spec", "Value"])
b_rows = [
    ("Nominal pipe diameter (in)", "6 [S6]"),
    ("Weight per linear foot (lb/ft)", "2.40 [S6]"),
    ("Material cost per linear foot (USD/ft)", "$5.10 [S7]"),
    ("Hazen-Williams roughness coefficient, C (PVC/plastic)", "150 [S3]"),
]
for r in b_rows:
    ws.append(r)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [45, 30])

# ---------- Appendix C ----------
ws = wb.create_sheet("Appendix C — Margin Main")
ws.append(["Spec", "Value"])
c_rows = [
    ("Nominal pipe diameter (in)", "8 [S6]"),
    ("Weight per linear foot (lb/ft)", "4.16 [S6]"),
    ("Material cost per linear foot (USD/ft)", "$7.30 [S8]"),
    ("Hazen-Williams roughness coefficient, C (PVC/plastic)", "150 [S3]"),
]
for r in c_rows:
    ws.append(r)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [45, 30])

# ---------- Appendix D ----------
ws = wb.create_sheet("Appendix D — Code 430")
ws.append(["Design criterion", "Value"])
d_rows = [
    ("Maximum allowable pipeline flow velocity at design capacity (ft/sec)", "5 [S1]"),
    ("Friction-loss method required for certification", "Hazen-Williams equation"),
    ("Hazen-Williams equation form", "h_f = 4.732 · (Q / C_HW)^1.85 · (L / D^4.87) [S2]"),
]
for r in d_rows:
    ws.append(r)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [45, 55])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- Appendix E ----------
ws = wb.create_sheet("Appendix E — EQIP Terms")
ws.append(["Term", "Value"])
e_rows = [
    ("Reimbursement basis", "Per pound of pipe installed"),
    ("Standard cost-share rate (% of eligible cost)", "75% [S4]"),
    ("Reimbursement rate per pound of pipe installed (USD/lb)", "$4.22 [S5]"),
    ("Condition of payment", "Reimbursement is forfeited in full if the installed design does not certify against NRCS Code 430"),
]
for r in e_rows:
    ws.append(r)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [35, 55])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- Appendix F ----------
ws = wb.create_sheet("Appendix F — Pump Energy")
ws.append(["Quantity", "Value"])
f_rows = [
    ("Pump power source (fuel type) for this job", "Electricity (grid-powered pump) [MA7]"),
    ("Energy-cost conversion factor", "Electricity: 0.885 whp-hr/kWh. Diesel: 12.5 whp-hr/gal; Gasoline: 8.55 whp-hr/gal [S9]"),
    ("Local price per unit of pump fuel or electricity (USD)", "$0.0917 per kWh [S12]"),
    ("Brake-horsepower sizing formula constants", "BHP = (GPM × TDH) ÷ (3,960 × E); E = 0.75 typical [S10]"),
]
for r in f_rows:
    ws.append(r)
style_header_row(ws, 1, 2)
set_body_font(ws)
col_widths(ws, [40, 60])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- Appendix G — Sources ----------
ws = wb.create_sheet("Appendix G — Sources")
ws.append(["#", "Claim", "Value", "Source", "Date", "Confidence"])
g_rows = [
    ("S1", "NRCS Code 430 maximum plastic-pipe design velocity", "5 ft/sec", "USDA NRCS, Conservation Practice Standard 430-CPS-2 (see case.json sources[] for URL)", "2024-11", "HIGH"),
    ("S2", "Hazen-Williams friction-loss equation", "h_f = 4.732(Q/C_HW)^1.85(L/D^4.87)", "USDA NRCS, NEH Part 634 Ch.4 eq. 4-28", "2022-02", "HIGH"),
    ("S3", "Hazen-Williams roughness coefficient for PVC/PE/HDPE pipe", "C = 150", "USDA NRCS, NEH Part 634 Ch.4 Fig. 4-8", "2022-02", "HIGH"),
    ("S4", "Standard EQIP cost-share payment rate", "up to 75%", "7 CFR § 1466.23(b)(1)(i)", "current", "HIGH"),
    ("S5", "EQIP Practice 430 reimbursement rate, PVC, standard rate", "$4.22/lb", "USDA NRCS Kansas FY2025 EQIP Payment Schedule", "2024-11", "HIGH"),
    ("S6", "PIP weight per linear foot, 6in & 8in nominal", "6in: 2.40 lb/ft; 8in: 4.16 lb/ft", "JM Eagle Irrigation P.I.P. catalog", "undated", "MEDIUM"),
    ("S7", "6in 125psi gasketed PVC pipe, price per foot", "$5.10/ft", "Direct Pivot Parts online catalog", "current listing", "MEDIUM"),
    ("S8", "8in 125psi gasketed PVC pipe, price per foot", "$7.30/ft", "Direct Pivot Parts online catalog", "current listing", "MEDIUM"),
    ("S9", "Water-horsepower-hours per unit of pumping energy, by fuel", "Electricity: 0.885 whp-hr/kWh", "UGA CAES Field Report C965 (Nebraska Pumping Plant Performance Criteria)", "reviewed 2012-06", "MEDIUM"),
    ("S10", "Brake-horsepower pump-sizing formula & typical pump efficiency", "BHP=(GPM×TDH)/(3960×E); E=0.75", "Texas A&M AgriLife Extension EB-6011", "2021-03", "HIGH"),
    ("S12", "U.S. average retail price of electricity, industrial sector", "9.17 cents/kWh", "U.S. EIA, Electricity Monthly Update", "2026-06", "HIGH"),
]
for r in g_rows:
    ws.append(r)
style_header_row(ws, 1, 6)
set_body_font(ws)
col_widths(ws, [6, 40, 24, 40, 16, 12])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- D1-D6 Work Area ----------
ws = wb.create_sheet("D1-D6 Work Area")
ws["A1"] = "Work Area — D1 through D6"
style_title(ws, "A1", size=13)
ws["A2"] = ("This sheet is intentionally blank except for column headers. Pull your raw inputs "
            "from the Appendix A/B/C/D/E/F sheets -- do not retype values by hand where you can "
            "reference the cell.")
ws["A2"].font = Font(name=ARIAL, italic=True)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:E2")
ws.row_dimensions[2].height = 40

sections = [
    (4, "D1 — Total-annualized-cost expression per option",
     ["Variable / constant / coefficient / term labeled", "Distributive-property expansion shown"]),
    (9, "D2 — Substitution table",
     ["Design flow rate substituted", "Run length substituted", "Diameter / C substituted", "Evaluated dollar figure"]),
    (14, "D3 — Simplified expression, like terms combined",
     ["Material-cost term", "Energy-cost term", "EQIP-reimbursement term", "Simplified total"]),
    (19, "D4 — Certification check (one-step equation)",
     ["Max flow rate before ceiling, Lean Main", "Max flow rate before ceiling, Margin Main", "Pass/fail at actual design flow"]),
    (24, "D5 — Rearranged Hazen-Williams for diameter",
     ["Multi-step rearrangement shown", "Stated energy-cost cap used", "Minimum diameter found"]),
    (29, "D6 — Break-even & recommendation",
     ["Equation set equal (D3 expressions)", "Break-even flow rate solved", "Actual flow rate vs. break-even", "Recommendation, defended with D4 and D6"]),
]
for start_row, title, headers in sections:
    ws.cell(row=start_row, column=1, value=title).font = Font(name=ARIAL, bold=True, size=12)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=start_row + 1, column=c, value=h)
    style_header_row(ws, start_row + 1, len(headers))
    for blank_r in range(start_row + 2, start_row + 4):
        for c in range(1, len(headers) + 1):
            ws.cell(row=blank_r, column=c, value=None)

col_widths(ws, [32, 32, 32, 32, 32])

wb.save("cases/the-mainline-margin/published/The Mainline Margin - Data.xlsx")
print("wrote The Mainline Margin - Data.xlsx")
