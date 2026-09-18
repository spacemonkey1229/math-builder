import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = "1A3A5C"
wb = openpyxl.Workbook()

def style_header(ws, row=1, ncols=6):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def set_font(ws, rows, cols, name="Arial", size=10):
    for r in rows:
        for c in cols:
            ws.cell(row=r, column=c).font = Font(name=name, size=size)

# ---------------- README ----------------
ws = wb.active
ws.title = "README"
ws["A1"] = "The Manifest Split — Data Workbook"
ws["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)
readme_rows = [
    ("", ""),
    ("Case", "The Manifest Split (cases/the-manifest-split)"),
    ("Purpose", "Raw, cited data for Marcus Feld's Combined-Lane-vs-Split-Lanes container "
                "inspection decision."),
    ("", ""),
    ("Sheet", "Maps to"),
    ("Appendix B — Flagged Manifest", "case.md Appendix B (gate-scan counts H, R, H∪R)"),
    ("Appendix C — Inspection Times", "case.md Appendix C (per-container check/re-handle times)"),
    ("Appendix D — Modeling Assumptions", "case.md Appendix D (MA1-MA6)"),
    ("Work Area", "Blank — for your own D1-D4 calculations. Nothing here is pre-computed."),
    ("", ""),
    ("Note", "This workbook gives you the DATA, not the ANSWER. The intersection count, the "
             "inspector-hour totals, the per-hour rates, the percent margin, and the Venn "
             "diagram are your own work — see Challenge 3 (D1-D4) in the case study."),
]
r = 2
for a, b in readme_rows:
    ws.cell(row=r, column=1, value=a).font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=r, column=2, value=b).font = Font(name="Arial", size=10)
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
    r += 1
autosize(ws, [26, 90])

# ---------------- Appendix B ----------------
ws = wb.create_sheet("Appendix B - Flagged Manifest")
ws.append(["Set", "Count", "How it's known", "Source (case.json sources[].id)"])
rows = [
    ["Hazmat-flagged containers (H)", 90, "Gate-scan hazmat placard/documentation trigger count, this sailing", "MA1"],
    ["Reefer-flagged containers (R)", 130, "Gate-scan reefer plug-request count, this sailing", "MA1"],
    ["Total distinct flagged containers (H union R)", 195, "Total unique container IDs that triggered either flag, read directly off the gate-scan log", "MA1"],
    ["Flagged for both (H intersect R)", "-- calculate in Work Area (D1) --", "Not given directly — derive via inclusion-exclusion", "--"],
]
for row in rows:
    ws.append(row)
style_header(ws, ncols=4)
set_font(ws, range(2, len(rows) + 2), range(1, 5))
autosize(ws, [40, 14, 60, 26])

# ---------------- Appendix C ----------------
ws = wb.create_sheet("Appendix C - Inspection Times")
ws.append(["Configuration", "Check", "Time per container (minutes)", "Source"])
rows = [
    ["The Combined Lane", "Hazmat check (generalist crew)", 8, "MA2"],
    ["The Combined Lane", "Reefer check (generalist crew)", 7, "MA3"],
    ["The Split Lanes", "Hazmat check (specialist crew)", 6, "MA4"],
    ["The Split Lanes", "Reefer check (specialist crew)", 5, "MA5"],
    ["The Split Lanes", "Re-handle (crane pull, yard move, re-queue) — dual-flagged containers only", 11, "MA6"],
]
for row in rows:
    ws.append(row)
style_header(ws, ncols=4)
set_font(ws, range(2, len(rows) + 2), range(1, 5))
autosize(ws, [22, 55, 16, 12])

# ---------------- Appendix D ----------------
ws = wb.create_sheet("Appendix D - Assumptions")
ws.append(["#", "Assumption", "Value used", "Basis"])
rows = [
    ["MA1", "Tonight's flagged-manifest counts (H, R, and H union R)", "H = 90, R = 130, H union R = 195",
     "No public source publishes a specific terminal's per-sailing hazmat/reefer flag breakdown — internal gate-scan data. Set at a scale consistent with a mid-sized terminal processing a multi-thousand-TEU vessel call."],
    ["MA2", "The Combined Lane hazmat check time", "8 minutes/container",
     "No public source; set higher than MA4 to reflect a generalist crew's task-switching cost."],
    ["MA3", "The Combined Lane reefer check time", "7 minutes/container",
     "Same basis and sensitivity class as MA2, applied to the reefer half of the combined check."],
    ["MA4", "The Split Lanes hazmat check time", "6 minutes/container",
     "Set below MA2 to reflect a specialist crew running one procedure repeatedly."],
    ["MA5", "The Split Lanes reefer check time", "5 minutes/container",
     "Same basis as MA4; consistent with, but not directly sourced from, S2's 'a five-minute inspection.'"],
    ["MA6", "Re-handle time for a dual-flagged container", "11 minutes/container",
     "No public source; set above S3's ~2-minute-per-crane-move floor to reflect a fuller re-handle process (un-stack + yard move + re-queue)."],
]
for row in rows:
    ws.append(row)
style_header(ws, ncols=4)
set_font(ws, range(2, len(rows) + 2), range(1, 5))
autosize(ws, [6, 42, 26, 60])

# ---------------- Work Area ----------------
ws = wb.create_sheet("Work Area")
ws["A1"] = "Work Area — your own calculations (nothing pre-filled)"
ws["A1"].font = Font(name="Arial", bold=True, size=12, color=NAVY)
labels = [
    "D1: H (hazmat-flagged) =",
    "D1: R (reefer-flagged) =",
    "D1: H union R (union, from gate-scan log) =",
    "D1: H intersect R (intersection) = H + R - (H union R) =",
    "D1: H-only = H - (H intersect R) =",
    "D1: R-only = R - (H intersect R) =",
    "D2: Combined Lane hazmat minutes = H x 8 =",
    "D2: Combined Lane reefer minutes = R x 7 =",
    "D2: Combined Lane total minutes =",
    "D2: Combined Lane total inspector-hours =",
    "D2: Split Lanes hazmat minutes = H x 6 =",
    "D2: Split Lanes reefer minutes = R x 5 =",
    "D2: Split Lanes re-handle minutes = (H intersect R) x 11 =",
    "D2: Split Lanes total minutes =",
    "D2: Split Lanes total inspector-hours =",
    "D3: Combined Lane rate = (H union R) / Combined Lane hours =",
    "D3: Split Lanes rate = (H union R) / Split Lanes hours =",
    "D3: Percent margin between the two rates =",
    "D3: Which configuration wins tonight? =",
    "D4: Written recommendation (see case study Challenge 3, D4) =",
]
r = 3
for lab in labels:
    ws.cell(row=r, column=1, value=lab).font = Font(name="Arial", size=10)
    ws.cell(row=r, column=2).font = Font(name="Arial", size=10)
    ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="FFF8E1")
    r += 1
autosize(ws, [56, 24])

wb.save("The Manifest Split - Data.xlsx")
print("wrote xlsx")
