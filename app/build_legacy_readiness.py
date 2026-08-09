#!/usr/bin/env python3
"""
build_legacy_readiness.py — a readiness scorecard for the 55 legacy corpus
cases, honestly scoped to what "readiness" can mean for a document written
before this pipeline's doctrine existed.

WHY THIS IS A SEPARATE SCRIPT, NOT A `tier1_structural.py --case <legacy-id>`:
tier1_structural.py operates on `cases/<id>/case.md` + `case.json` — a shape
that legacy cases do not have (they are raw .docx files in source/, never
run through case-write). Two real samples read before writing this script
(A Coach's Dilemma's case doc and Guide Notes) confirmed the deeper reason
a literal port would mislead more than it would help:
  - The current CHALLENGE_BOILERPLATE.md text (250+ words, a research-task
    Challenge 1) postdates this corpus. Legacy Challenge sections use older,
    real wording ("Post an 'excellent' write-up of at least 500 words...",
    no research-task Challenge 1 at all). A verbatim diff would fail on
    every single legacy case, not because they are bad, but because the
    doctrine they'd be diffed against didn't exist yet when they were
    written.
  - Legacy Guide Notes use a real but different facilitation shape (Open /
    Q1 / Q2 / Q3 / Close / Lessons Learned / Standards Touched On), not the
    11-heading Intro/Launch/Anchor-1..4/Data-Deep-Dive/Twist/Close/Lessons-
    Learned/Standards-Touched-On ladder SPEC §7.1 requires. SPEC's own text
    says the ladder was "derived from 39 Guide Notes documents" — i.e. it
    is the median PATTERN across them, not a heading vocabulary any one of
    them actually used.

So this script checks two genuinely different things and does not conflate
them:
  (A) STRUCTURAL SHAPE — do the core story beats SPEC was built from exist
      at all (Overview/Background/binary Dilemma/Challenges/Appendices),
      loosely matched by heading text, not exact wording. This is a fair
      question to ask of a pre-doctrine document.
  (B) CURRENT-DOCTRINE ARTIFACTS — the newer, pipeline-era requirements
      (verbatim boilerplate, the 11-anchor ladder by name, a named Twist,
      a timed 45-minute plan, cohort bands, interactive HTML, a hover-
      hyperlink glossary). These are known, near-universally absent by
      construction — not a defect discovered per-case, a fact about when
      the doctrine was written relative to when the corpus was. Reported
      as a flat count, not a per-case shame list.

Read-only over source/*.docx and readmodel/*.json. Writes only
readmodel/legacy_readiness.json.

Usage:
  python3 app/build_legacy_readiness.py
  python3 app/build_legacy_readiness.py --verify
"""
import argparse, datetime, json, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

try:
    import docx
except ImportError:
    print("python-docx not installed. pip install python-docx", file=sys.stderr)
    sys.exit(1)


def load(name):
    return json.load(open(os.path.join(ROOT, "readmodel", name), encoding="utf-8"))


def doc_text(path):
    try:
        d = docx.Document(path)
    except Exception as e:
        return None, str(e)
    paras = [p.text for p in d.paragraphs]
    # also pull table cell text -- several legacy docs put Challenge/Appendix
    # content in tables, not paragraphs
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                paras.extend(p.text for p in cell.paragraphs)
    return paras, None


SECTION_PAT = {
    "overview": re.compile(r"^\s*overview\s*:?\s*$", re.I),
    "background": re.compile(r"^\s*background\s*:?\s*$", re.I),
    "dilemma": re.compile(r"^\s*(the\s+)?dilemma\s*:?\s*$", re.I),
    "challenges": re.compile(r"^\s*challenges?\s*:?\s*$", re.I),
}
APPENDIX_PAT = re.compile(r"^\s*appendix\s+[a-z]\b", re.I)
BINARY_A_PAT = re.compile(r"^\s*a\.?\)", re.I)
BINARY_B_PAT = re.compile(r"^\s*b\.?\)", re.I)
BINARY_C_PAT = re.compile(r"^\s*c\.?\)", re.I)
TITLE_PAT = re.compile(r"case study\s*:", re.I)

ANCHOR_HEADINGS = ["intro", "launch", "anchor 1", "anchor 2", "anchor 3", "anchor 4",
                    "data deep dive", "twist", "close", "lessons learned", "standards touched on"]


def check_case_doc(paras):
    text_all = "\n".join(paras).lower()
    found_sections = {k: any(p_pat.match(p) for p in paras) for k, p_pat in SECTION_PAT.items()}
    n_appendices = len({p.strip().split(":")[0].split()[1].rstrip(".,")
                         for p in paras if APPENDIX_PAT.match(p.strip())}
                        ) if any(APPENDIX_PAT.match(p.strip()) for p in paras) else 0
    has_a = any(BINARY_A_PAT.match(p.strip()) for p in paras)
    has_b = any(BINARY_B_PAT.match(p.strip()) for p in paras)
    has_c = any(BINARY_C_PAT.match(p.strip()) for p in paras)
    return {
        "has_title_block": any(TITLE_PAT.search(p) for p in paras[:3]),
        "has_overview": found_sections["overview"],
        "has_background": found_sections["background"],
        "has_dilemma_heading": found_sections["dilemma"],
        "has_binary_options": has_a and has_b,
        "has_third_option": has_c,
        "has_challenges_heading": found_sections["challenges"],
        "n_appendices_detected": n_appendices,
        "mentions_how_to_research": "how to" in text_all and ("research" in text_all or "questions to research" in text_all),
    }


def check_guide_notes(paras):
    text_all = "\n".join(paras).lower()
    headings_found = [h for h in ANCHOR_HEADINGS if any(p.strip().lower().rstrip(":").strip() == h for p in paras)]
    has_timing = bool(re.search(r"\b\d{1,2}\s*(?:minutes?|min)\b", text_all)) or bool(re.search(r"\b\d{1,2}:\d{2}\b", text_all))
    return {
        "anchor_headings_matched": headings_found,
        "n_anchor_headings_matched": len(headings_found),
        "mentions_twist": "twist" in text_all,
        "mentions_vote": "vote" in text_all,
        "has_some_timing_language": has_timing,
        "has_lessons_learned": any(p.strip().lower().rstrip(":") == "lessons learned" for p in paras),
        "has_standards_touched_on": any("standards touched" in p.strip().lower() for p in paras),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--verify", action="store_true")
    a = ap.parse_args()

    ci = load("corpus_index.json")
    spine = load("standards_spine.json")
    # corpus_index.json titles escape an apostrophe as "_s" in folder-derived titles
    # (e.g. "A Coach_s Dilemma"); standards_spine.json.cases[] carries the real
    # apostrophe. Same normalization build_dashboard.py's reconcile() already uses
    # for exactly this reason -- without it, a genuinely-tagged case undercounts.
    def norm_apos(s):
        return s.replace("_s", "'s")
    tagged_titles = set(spine.get("cases", []))
    tagged_titles_norm = {norm_apos(t) for t in tagged_titles}

    legacy = [c for c in ci["cases"] if c["status"] == "legacy"]
    rows = []
    errors = []
    for c in legacy:
        base = os.path.join(ROOT, c["path"])
        case_files = c["assets"].get("case_doc") or []
        row = {"title": c["title"], "path": c["path"], "has_guide": c["has_guide"],
               "has_socratic_plan": c["has_socratic_plan"],
               "standards_tagged": c["title"] in tagged_titles or norm_apos(c["title"]) in tagged_titles_norm}
        if not case_files:
            row["case_doc_error"] = "no case_doc asset recorded"
            rows.append(row)
            continue
        paras, err = doc_text(os.path.join(base, case_files[0]))
        if err:
            row["case_doc_error"] = err
        else:
            row["structure"] = check_case_doc(paras)

        gn = c["assets"].get("guide_notes") or []
        if gn:
            gparas, gerr = doc_text(os.path.join(base, gn[0]))
            if gerr:
                row["guide_doc_error"] = gerr
            else:
                row["facilitation"] = check_guide_notes(gparas)
        rows.append(row)

    ok_rows = [r for r in rows if "structure" in r]
    n = len(ok_rows)

    def pct(pred):
        k = sum(1 for r in ok_rows if pred(r))
        return k, round(100.0 * k / n, 1) if n else 0.0

    struct_summary = {}
    for key in ("has_title_block", "has_overview", "has_background", "has_dilemma_heading",
                "has_binary_options", "has_challenges_heading", "mentions_how_to_research"):
        k, p = pct(lambda r, key=key: r["structure"].get(key))
        struct_summary[key] = {"n": k, "pct": p}
    k, p = pct(lambda r: r["structure"].get("has_third_option"))
    struct_summary["has_third_option (should be 0 -- SPEC requires strictly binary)"] = {"n": k, "pct": p}
    struct_summary["n_appendices_detected_median"] = (
        sorted(r["structure"]["n_appendices_detected"] for r in ok_rows)[n // 2] if n else None
    )

    fac_rows = [r for r in rows if "facilitation" in r]
    nf = len(fac_rows)

    def fpct(pred):
        k = sum(1 for r in fac_rows if pred(r))
        return k, round(100.0 * k / nf, 1) if nf else 0.0

    fac_summary = {}
    for key in ("mentions_twist", "mentions_vote", "has_some_timing_language",
                "has_lessons_learned", "has_standards_touched_on"):
        k, p = fpct(lambda r, key=key: r["facilitation"].get(key))
        fac_summary[key] = {"n": k, "pct": p}
    anchor_dist = {}
    for r in fac_rows:
        k = r["facilitation"]["n_anchor_headings_matched"]
        anchor_dist[str(k)] = anchor_dist.get(str(k), 0) + 1
    fac_summary["anchor_headings_matched_distribution (of 11 possible, exact-heading match only)"] = anchor_dist

    out = {
        "_schema": {
            "_comment": "DERIVED — regenerate with app/build_legacy_readiness.py. "
                        "Reads source/**/*.docx directly, read-only. "
                        "Never hand-edit. See the module docstring for why this checks "
                        "STRUCTURAL SHAPE (A) separately from CURRENT-DOCTRINE ARTIFACTS (B) "
                        "instead of running tier1_structural.py's exact checks against legacy docs.",
            "scope_note": "Legacy cases are NOT run through case-write/case-research/case-readiness "
                           "and are not held to the CURRENT locked CHALLENGE_BOILERPLATE.md text or "
                           "the exact 11-anchor heading vocabulary, both of which postdate this corpus. "
                           "This scorecard checks structural shape (loose heading match) and honestly "
                           "flags what's absent by construction, not by individual-case defect.",
        },
        "_generated_at": None,
        "n_legacy_cases": len(legacy),
        "n_case_docs_readable": n,
        "n_case_docs_errored": len(legacy) - n,
        "n_with_guide_notes_readable": nf,
        "structural_shape_summary": struct_summary,
        "facilitation_summary": fac_summary,
        "known_absent_by_construction": {
            "cohort_bands (MCS-25)": "0 of 55 -- introduced with Cold Chain's rewrite, 2026-07-31",
            "interactive_html (MCS-37)": "0 of 55 -- introduced with Cold Chain, 2026-07-31",
            "hover_hyperlink_glossary (MCS-39)": "0 of 55 -- introduced with Cold Chain, 2026-08-01",
            "verbatim_current_boilerplate": "effectively 0 of 55 -- current CHALLENGE_BOILERPLATE.md wording postdates this corpus",
        },
        "standards_tagged": sum(1 for r in rows if r.get("standards_tagged")),
        "cases": rows,
    }

    outpath = os.path.join(ROOT, "readmodel", "legacy_readiness.json")
    if a.verify:
        if not os.path.exists(outpath):
            print("MISSING readmodel/legacy_readiness.json — run without --verify first")
            sys.exit(1)
        existing = json.load(open(outpath, encoding="utf-8"))
        e2 = dict(existing); e2.pop("_generated_at", None)
        o2 = dict(out); o2.pop("_generated_at", None)
        if e2 != o2:
            print("STALE — does not match a fresh derivation")
            sys.exit(1)
        print("legacy_readiness.json is current")
        return

    out["_generated_at"] = datetime.date.today().isoformat()
    json.dump(out, open(outpath, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print(f"wrote {outpath}")
    print(f"read {n}/{len(legacy)} case docs, {nf} guide notes docs")
    print(json.dumps(struct_summary, indent=1))
    print(json.dumps(fac_summary, indent=1))


if __name__ == "__main__":
    main()
